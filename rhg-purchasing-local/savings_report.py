"""Price savings report: the background read, the PDF, the Excel file and the email.

The screen (Purchasing → Price savings) asks for any date range with filters; the
months it needs are read once in the background and the report itself is worked out
from the local file in well under a second. The same report goes out as an email with
the PDF and the Excel file attached: on demand from the screen, and on the first working
day of each month for the month before, to the Purchasing Manager and Coordinator set in
Settings → Email notifications.
"""
import html
import io
import json
import threading
import time
import traceback
from datetime import date, datetime, timedelta

import notify
import savings
import smtp_config
import store

STATE_FILE = store.DATA_DIR / "savings_report.json"
MAX_DAYS = 731
_jobs, _jobs_lock = {}, threading.Lock()
_state_lock = threading.Lock()


# ---------------------------------------------------------------- the request
def parse_range(date_from, date_to, today=None):
    today = today or date.today()
    if not date_from or not date_to:
        first = today.replace(day=1)
        b = first - timedelta(days=1)
        return b.replace(day=1), b                      # default: the last complete month
    a, b = date.fromisoformat(date_from[:10]), date.fromisoformat(date_to[:10])
    b = min(b, today)
    if a > b:
        raise ValueError("The start date is after the end date.")
    if (b - a).days > MAX_DAYS:
        raise ValueError("Choose a range of two years or less.")
    return a, b


def _start_load(client, company, months, force):
    with _jobs_lock:
        job = _jobs.get(company)
        if job and job["state"] == "running":
            return job
        job = _jobs[company] = {"state": "running", "done": 0, "total": len(months), "errors": [],
                                "started": time.time()}

    def run():
        def tick():
            with _jobs_lock:
                job["done"] += 1
        errors = savings.ensure(client, company, months, force=force, tick=tick)
        with _jobs_lock:
            job.update(state="done", errors=errors, finished=time.time())
    threading.Thread(target=run, daemon=True, name="savings-load").start()
    return job


def request(client, company, date_from, date_to, filters, refresh=False):
    a, b = parse_range(date_from, date_to)
    months = savings.needed_months(a, b)
    missing = savings.missing_months(company, months)
    if refresh:
        recent = [m for m in months if m >= (date.today() - timedelta(days=savings.OPEN_DAYS)).strftime("%Y-%m")]
        missing = sorted(set(missing) | set(recent))
    with _jobs_lock:
        job = dict(_jobs.get(company) or {})
    if job.get("state") == "running" or missing:
        if job.get("state") != "running":
            job = dict(_start_load(client, company, sorted(missing, reverse=True), refresh))
        return {"pending": True, "progress": {"done": job.get("done", 0), "total": job.get("total", 0)},
                "period": {"from": a.isoformat(), "to": b.isoformat()},
                "note": "Reading the purchase invoice lines of these months from Business Central. A month "
                        "is read once; the next report on the same months is immediate."}
    out = savings.report(client, company, a, b, **filters)
    out["errors"] = job.get("errors") or []
    out["company"] = company
    return out


# ---------------------------------------------------------------- formatting helpers
def _m(v):
    return "–" if v is None else f"{v:,.0f}"


def _filters_text(r):
    f = r["filters"]
    parts = []
    if f.get("location"):
        parts.append(f"location {f['location']}")
    if f.get("vendor"):
        name = next((v["name"] for v in r["options"]["vendors"] if v["no"] == f["vendor"]), f["vendor"])
        parts.append(f"supplier {name}")
    if f.get("category"):
        parts.append(f"category {f['category']}")
    if f.get("stream"):
        parts.append(f"type {f['stream']}")
    if f.get("q"):
        parts.append(f"items matching '{f['q']}'")
    return ", ".join(parts) or "all locations, suppliers and items"


def _period_text(r):
    p, l = r["period"], r["last_year"]
    return f"{p['from']} to {p['to']} against {l['from']} to {l['to']}"


# ---------------------------------------------------------------- PDF
def render_pdf(r, currency):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy, ink, muted = colors.HexColor("#08303f"), colors.HexColor("#0e2733"), colors.HexColor("#5c7482")
    sky, line = colors.HexColor("#e4f3f9"), colors.HexColor("#d9dcdd")
    green, red = colors.HexColor("#0e8a57"), colors.HexColor("#c0392b")
    st = {"h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=16, textColor=ink, leading=20),
          "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11.5, textColor=ink, leading=15,
                               spaceBefore=10, spaceAfter=5),
          "p": ParagraphStyle("p", fontName="Helvetica", fontSize=8.5, textColor=muted, leading=11),
          "c": ParagraphStyle("c", fontName="Helvetica", fontSize=7.8, textColor=ink, leading=9.5)}
    s = r["summary"]
    e = lambda t: html.escape(str(t or ""))

    def grid(rows, widths, num_cols=(), color_col=None):
        t = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        style = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.8), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.8),
                 ("TEXTCOLOR", (0, 0), (-1, -1), ink), ("BACKGROUND", (0, 0), (-1, 0), sky),
                 ("LINEBELOW", (0, 0), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        for c in num_cols:
            style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
        if color_col is not None:
            for i, row in enumerate(rows[1:], start=1):
                v = row[color_col]
                if isinstance(v, str) and v.startswith("-"):
                    style.append(("TEXTCOLOR", (color_col, i), (color_col, i), green))
                elif isinstance(v, str) and v[:1].isdigit() and v not in ("0",):
                    style.append(("TEXTCOLOR", (color_col, i), (color_col, i), red))
        t.setStyle(TableStyle(style))
        return t

    story = [Paragraph("Price savings against last year", st["h1"]),
             Paragraph(f"{e(r.get('company'))} · {e(_period_text(r))} · {e(_filters_text(r))} · amounts in "
                       f"{e(currency)} excluding VAT", st["p"]), Spacer(1, 8)]
    story.append(grid([
        ["Last year's spend", "Savings (cheaper)", "Extra (dearer)", "Net price effect", "Volume and mix",
         "This period's spend", "Items compared"],
        [_m(s["ly"]), _m(s["save"]), _m(s["extra"]), _m(s["net"]), _m(s["volume"]), _m(s["cur"]),
         f"{s['items']} ({_m(s['matched'])} {currency})"]],
        [36 * mm] * 6 + [50 * mm], num_cols=range(0, 7)))
    story.append(Paragraph("A negative net price effect is a saving. " + e(r["definitions"]["method"]), st["p"]))

    if len(r["months"]) > 1:
        story.append(Paragraph("By month", st["h2"]))
        story.append(grid([["Month", "Last year", "This year", "Savings", "Extra", "Net price effect", "Volume and mix"]]
                          + [[m["month"], _m(m["ly"]), _m(m["cur"]), _m(m["save"]), _m(m["extra"]), _m(m["net"]),
                              _m(m["volume"])] for m in r["months"]],
                          [26 * mm] + [34 * mm] * 6, num_cols=range(1, 7), color_col=5))

    def group_table(title, rows_in, label, name_of=lambda g: g["key"]):
        story.append(Paragraph(title, st["h2"]))
        story.append(grid([[label, "Last year", "This year", "Savings", "Extra", "Net price effect", "Items"]]
                          + [[Paragraph(e(name_of(g)), st["c"]), _m(g["ly"]), _m(g["cur"]), _m(g["save"]),
                              _m(g["extra"]), _m(g["net"]), str(g["items"])] for g in rows_in],
                          [78 * mm, 30 * mm, 30 * mm, 30 * mm, 30 * mm, 32 * mm, 18 * mm], num_cols=range(1, 7),
                          color_col=5))
    group_table("By purchase type", r["streams"], "Type")
    group_table("By supplier", r["vendors"], "Supplier", lambda g: f"{g['name']} ({g['key']})")
    group_table("By item category", r["categories"], "Category")
    group_table("By location", r["locations"], "Location", lambda g: g["key"] or "(no location)")

    story.append(Paragraph(f"Every item ({r['items_total']})", st["h2"]))
    rows_ = [["Item", "Category", "Status", "Unit cost LY", "Unit cost now", "Change", "Quantity", "Spend now",
              "Price effect", "Main suppliers"]]
    status_label = {"both": "compared", "new": "new", "dropped": "not bought", "excluded": "left out", "other": ""}
    for it in r["items"]:
        rows_.append([Paragraph(e(f"{it['desc']} ({it['item']})"), st["c"]), e(it["category"]),
                      status_label.get(it["status"], it["status"]),
                      f"{it['ly_avg']:,.2f}" if it["ly_avg"] else "–", f"{it['cur_avg']:,.2f}" if it["cur_avg"] else "–",
                      f"{it['pct']:+.1f}%" if it["pct"] is not None else "",
                      f"{it['cur_qty']:,.1f}", _m(it["cur_cost"]), _m(it["effect"]) if it["effect"] is not None else "",
                      Paragraph(e(", ".join(it["vendors"])), st["c"])])
    story.append(grid(rows_, [62 * mm, 18 * mm, 17 * mm, 21 * mm, 21 * mm, 15 * mm, 20 * mm, 24 * mm, 22 * mm, 53 * mm],
                      num_cols=range(3, 9), color_col=8))
    story.append(Spacer(1, 6))
    for k in ("source", "volume", "excluded", "suppliers"):
        story.append(Paragraph(e(r["definitions"][k]), st["p"]))

    def frame(canvas, doc):
        canvas.saveState()
        w, h = landscape(A4)
        canvas.setFillColor(navy)
        canvas.rect(0, h - 9 * mm, w, 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawString(12 * mm, h - 6 * mm, f"RHG PURCHASING · PRICE SAVINGS · {r.get('company', '')} · "
                                               f"{r['period']['from']} to {r['period']['to']}")
        canvas.setFillColor(muted)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(12 * mm, 7 * mm, "RHG Control Towers · read from Business Central, nothing changed there")
        canvas.drawRightString(w - 12 * mm, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=15 * mm,
                      bottomMargin=14 * mm, title=f"Price savings {r.get('company', '')} {r['period']['from']} "
                                                  f"{r['period']['to']}", author="RHG Control Towers"
                      ).build(story, onFirstPage=frame, onLaterPages=frame)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def render_xlsx(r, currency):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="E4F3F9"), Font(bold=True, color="0E2733")
    money, unit, pct = '#,##0;[Red]-#,##0', '#,##0.00', '0.0%'

    def sheet(ws, header, rows_in, formats, widths):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows_in:
            ws.append(row)
        for i, f in enumerate(formats, start=1):
            if f:
                for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell[0].number_format = f
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows_in:
            ws.auto_filter.ref = ws.dimensions

    s = r["summary"]
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Price savings against last year"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in (("Company", r.get("company")), ("Period", f"{r['period']['from']} to {r['period']['to']}"),
                 ("Compared with", f"{r['last_year']['from']} to {r['last_year']['to']}"),
                 ("Filters", _filters_text(r)), ("Currency", f"{currency}, excluding VAT"), ("", ""),
                 ("Last year's spend", s["ly"]), ("Savings from lower prices", -s["save"]),
                 ("Extra spending from higher prices", s["extra"]), ("Net price effect", s["net"]),
                 ("Volume and mix", s["volume"]), ("This period's spend", s["cur"]),
                 ("Items compared", s["items"]), ("Spend on compared items", s["matched"]),
                 ("Items left out (unit cost moved over 4x)", s["excluded"])):
        ws.append([k, v])
        if isinstance(v, (int, float)) and k != "Items compared":
            ws.cell(ws.max_row, 2).number_format = money
    ws.append([])
    for k in ("method", "source", "volume", "excluded", "suppliers"):
        ws.append([r["definitions"][k]])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22

    status_label = {"both": "compared", "new": "new this period", "dropped": "not bought this period",
                    "excluded": "left out", "other": ""}
    sheet(wb.create_sheet("Items"),
          ["Item", "Description", "Category", "Purchase type", "Status", "Qty last year", "Spend last year",
           "Unit cost last year", "Qty this period", "Spend this period", "Unit cost this period", "Change %",
           "Price effect (negative = saving)", "Main suppliers"],
          [[i["item"], i["desc"], i["category"], i["stream"], status_label.get(i["status"], i["status"]),
            i["ly_qty"], i["ly_cost"], i["ly_avg"], i["cur_qty"], i["cur_cost"], i["cur_avg"],
            (i["pct"] / 100) if i["pct"] is not None else None, i["effect"], ", ".join(i["vendors"])]
           for i in r["items"]],
          [None, None, None, None, None, unit, money, unit, unit, money, unit, pct, money, None],
          [12, 40, 12, 18, 16, 12, 15, 14, 12, 15, 14, 10, 18, 40])

    def group_sheet(name, rows_in, label, with_name=False):
        sheet(wb.create_sheet(name),
              [label] + (["Name"] if with_name else []) + ["Spend last year", "Spend this period", "Savings",
                                                            "Extra spending", "Net price effect", "Items compared"],
              [[g["key"]] + ([g.get("name")] if with_name else []) + [g["ly"], g["cur"], g["save"], g["extra"],
                                                                     g["net"], g["items"]] for g in rows_in],
              [None] + ([None] if with_name else []) + [money] * 5 + [None],
              [18] + ([40] if with_name else []) + [16] * 6)
    group_sheet("Suppliers", r["vendors"], "Supplier", True)
    group_sheet("Purchase types", r["streams"], "Type")
    group_sheet("Categories", r["categories"], "Category")
    group_sheet("Locations", r["locations"], "Location")
    sheet(wb.create_sheet("Months"),
          ["Month", "From", "To", "Spend last year", "Spend this year", "Savings", "Extra spending",
           "Net price effect", "Volume and mix"],
          [[m["month"], m["from"], m["to"], m["ly"], m["cur"], m["save"], m["extra"], m["net"], m["volume"]]
           for m in r["months"]],
          [None, None, None] + [money] * 6, [10, 12, 12] + [16] * 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_stem(r):
    f = r["filters"]
    tag = "_".join(x for x in (f.get("location"), f.get("vendor"), f.get("category")) if x)
    return (f"RHG_{(r.get('company') or '').replace(' ', '_')}_price_savings_{r['period']['from']}_"
            f"{r['period']['to']}" + (f"_{tag.replace(' ', '_')}" if tag else ""))


# ---------------------------------------------------------------- email
def render_email(r, currency, note=""):
    e = html.escape
    s = r["summary"]
    navy, primary, ink, muted, line, sky = "#08303f", "#0c475e", "#0e2733", "#5c7482", "#d9dcdd", "#e4f3f9"
    green, red = "#0e8a57", "#c0392b"
    public = (smtp_config.current().get("public_url") or "").rstrip("/")
    link = f"{public}/purchasing/#savings" if public else ""
    th = f"text-align:left;padding:7px 9px;background:{sky};color:{ink};font-size:12px;font-weight:600"
    thn = th.replace("text-align:left", "text-align:right")
    td = f"padding:7px 9px;border-bottom:1px solid #f0f1f2;font-size:12px;color:{ink}"
    tdn = td + ";text-align:right;white-space:nowrap"

    def kpi(label, value, sub, color=ink):
        return (f"<td style='padding:12px 14px;background:#fff;border:1px solid {line};border-radius:10px;"
                f"vertical-align:top;width:25%'><div style='font-size:12px;color:{muted}'>{e(label)}</div>"
                f"<div style='font-size:20px;font-weight:700;color:{color};margin:4px 0'>{e(value)}</div>"
                f"<div style='font-size:11px;color:{muted}'>{e(sub)}</div></td>")
    net_good = s["net"] <= 0
    kpis = ("<table role='presentation' cellspacing='8' cellpadding='0' style='width:100%;margin:0 -8px'><tr>"
            + kpi("Savings from lower prices", _m(s["save"]), "items now bought cheaper", green)
            + kpi("Extra from higher prices", _m(s["extra"]), "items now bought dearer", red)
            + kpi("Net saving on price" if net_good else "Net extra on price", _m(abs(s["net"])),
                  f"{abs(s['net']) / s['ly'] * 100:.1f}% of last year's spend" if s["ly"] else "",
                  green if net_good else red)
            + kpi("Volume and mix", _m(s["volume"]), f"spend {_m(s['ly'])} → {_m(s['cur'])}")
            + "</tr></table>")

    def items_table(rows_in, title, good):
        if not rows_in:
            return ""
        return (f"<h3 style='font-size:14px;color:{ink};margin:20px 0 8px'>{e(title)}</h3>"
                f"<table role='presentation' cellspacing='0' cellpadding='0' style='width:100%;border-collapse:collapse;"
                f"border:1px solid {line}'><tr><th style='{th}'>Item</th><th style='{thn}'>Last year</th>"
                f"<th style='{thn}'>Now</th><th style='{thn}'>Change</th><th style='{thn}'>{'Saved' if good else 'Extra'}</th>"
                f"<th style='{th}'>Supplier</th></tr>"
                + "".join(f"<tr><td style='{td}'>{e(i['desc'] or i['item'])} <span style='color:{muted}'>{e(i['item'])}</span></td>"
                          f"<td style='{tdn}'>{i['ly_avg']:,.2f}</td><td style='{tdn}'>{i['cur_avg']:,.2f}</td>"
                          f"<td style='{tdn};color:{green if good else red}'>{i['pct']:+.1f}%</td>"
                          f"<td style='{tdn}'><b>{_m(abs(i['effect']))}</b></td>"
                          f"<td style='{td};color:{muted}'>{e(', '.join(i['vendors'][:2]))}</td></tr>" for i in rows_in)
                + "</table>")
    compared = [i for i in r["items"] if i["status"] == "both"]
    best = sorted([i for i in compared if i["effect"] < 0], key=lambda i: i["effect"])[:10]
    worst = sorted([i for i in compared if i["effect"] > 0], key=lambda i: -i["effect"])[:10]
    vend = sorted(r["vendors"], key=lambda g: g["net"])
    vend_rows = vend[:5] + [g for g in reversed(vend[-5:]) if g not in vend[:5]]
    vend_table = ("" if not vend_rows else
                  f"<h3 style='font-size:14px;color:{ink};margin:20px 0 8px'>Suppliers with the largest price effect</h3>"
                  f"<table role='presentation' cellspacing='0' cellpadding='0' style='width:100%;border-collapse:collapse;"
                  f"border:1px solid {line}'><tr><th style='{th}'>Supplier</th><th style='{thn}'>Spend now</th>"
                  f"<th style='{thn}'>Savings</th><th style='{thn}'>Extra</th><th style='{thn}'>Net</th></tr>"
                  + "".join(f"<tr><td style='{td}'>{e(g['name'])}</td><td style='{tdn}'>{_m(g['cur'])}</td>"
                            f"<td style='{tdn};color:{green}'>{_m(g['save'])}</td><td style='{tdn};color:{red}'>{_m(g['extra'])}</td>"
                            f"<td style='{tdn}'><b>{_m(g['net'])}</b></td></tr>" for g in vend_rows) + "</table>")
    button = (f"<p style='margin:22px 0 0'><a href='{e(link)}' style='background:{primary};color:#fff;text-decoration:none;"
              f"padding:10px 18px;border-radius:8px;font-size:13px;font-weight:600;display:inline-block'>"
              f"Open Price savings</a></p>") if link else ""
    body = f"""<!doctype html><html><head><meta charset="utf-8"></head><body style="margin:0;background:#ecebeb;font-family:Figtree,Segoe UI,Arial,sans-serif;color:{ink}">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="background:#ecebeb"><tr><td align="center" style="padding:24px 12px">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="max-width:860px">
<tr><td style="background:{navy};color:#fff;padding:18px 22px;border-radius:12px 12px 0 0">
<div style="font-size:11px;letter-spacing:.12em;color:#a0d8ea">REVO HOSPITALITY GROUP · RHG PURCHASING</div>
<div style="font-size:20px;font-weight:700;margin-top:4px">Price savings against last year</div>
<div style="font-size:12px;color:#cfe3ea;margin-top:4px">{e(r.get('company'))} · {e(_period_text(r))} · {e(_filters_text(r))}</div></td></tr>
<tr><td style="background:#fff;padding:22px;border-radius:0 0 12px 12px">
{f"<div style='background:#fdf3e2;color:#b26a00;padding:8px 12px;border-radius:8px;font-size:12px;margin-bottom:14px'>{e(note)}</div>" if note else ""}
{kpis}
<p style="color:{muted};font-size:12px;margin:14px 0 0">{e(r['definitions']['method'])} Amounts in {e(currency)} excluding VAT. <b>The attached PDF and Excel file list every item, supplier, category and location.</b></p>
{items_table(best, "Largest savings", True)}{items_table(worst, "Largest extra spending", False)}{vend_table}{button}
<p style="color:{muted};font-size:11px;margin:24px 0 0;border-top:1px solid {line};padding-top:12px">Sent by RHG Control Towers from the posted purchase invoices in Business Central; nothing was changed there.</p>
</td></tr></table></td></tr></table></body></html>"""
    text = (f"RHG Purchasing · Price savings against last year · {r.get('company')}\n{_period_text(r)} · {_filters_text(r)}\n\n"
            f"Savings from lower prices: {_m(s['save'])}\nExtra from higher prices: {_m(s['extra'])}\n"
            f"Net price effect: {_m(s['net'])} (negative = saving)\nVolume and mix: {_m(s['volume'])}\n"
            f"Spend: {_m(s['ly'])} last year -> {_m(s['cur'])} now ({currency}, excl. VAT)\n\n"
            "The attached PDF and Excel file list every item." + (f"\n\nOpen: {link}" if link else ""))
    return body, text


def subject(r, test=False):
    s = r["summary"]
    word = "net saving" if s["net"] <= 0 else "net extra spending"
    return (("[Test] " if test else "") + f"[RHG Purchasing] {r.get('company')} · Price savings "
            f"{r['period']['from']} to {r['period']['to']} · {word} {_m(abs(s['net']))}")


def recipients(settings=None):
    n = (settings or store.get_settings()).get("notifications") or {}
    to = notify.parse_emails(n.get("manager", ""))[0] + notify.parse_emails(n.get("coordinator", ""))[0]
    cc = [c for c in notify.parse_emails(n.get("cost_controller", ""))[0] if c not in to]
    return list(dict.fromkeys(to)), cc


def send(r, currency, to=None, cc=None, actor="", test=False, note=""):
    if to is None:
        to, cc = recipients()
    if not to:
        raise RuntimeError("No Purchasing Manager or Coordinator email is set in Settings → Email notifications.")
    body, text = render_email(r, currency, note)
    stem = file_stem(r)
    files = [(f"{stem}.pdf", "application", "pdf", render_pdf(r, currency)),
             (f"{stem}.xlsx", "application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              render_xlsx(r, currency))]
    notify.send_mail(to, cc or [], subject(r, test), body, text, files)
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), "actor": actor, "to": to, "cc": cc or [],
             "period": r["period"], "filters": r["filters"], "net": r["summary"]["net"], "test": test}
    _log(entry)
    store.audit("email", "price savings report", actor or "scheduler", {}, entry, subject(r, test))
    return entry


# ---------------------------------------------------------------- monthly email
def _state():
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_state(s):
    tmp = STATE_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(s, indent=1, default=str), encoding="utf-8")
    tmp.replace(STATE_FILE)


def _log(entry):
    with _state_lock:
        s = _state()
        s["runs"] = (s.get("runs", []) + [entry])[-100:]
        _save_state(s)


def recent_runs(limit=10):
    with _state_lock:
        return list(reversed(_state().get("runs", [])))[:limit]


def first_working_day(y, m):
    d = date(y, m, 1)
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def next_monthly(now=None):
    n = store.get_settings().get("notifications") or {}
    if not n.get("savings_monthly"):
        return None
    now = now or datetime.now()
    hh, _, mm = (n.get("time") or "08:00").partition(":")
    for k in range(3):
        y, m = now.year + (now.month - 1 + k) // 12, (now.month - 1 + k) % 12 + 1
        d = first_working_day(y, m)
        t = datetime(d.year, d.month, d.day, int(hh), int(mm or 0))
        with _state_lock:
            done = _state().get("monthly_done")
        prev = (d.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
        if t + timedelta(hours=10) >= now and done != prev:
            return t.isoformat(timespec="minutes")
    return None


class Monthly:
    """Sends last month's report on the first working day of the month."""

    def __init__(self, client, companies_of, currency_of):
        self.client, self.companies_of, self.currency_of = client, companies_of, currency_of

    def due(self, now):
        n = store.get_settings().get("notifications") or {}
        if not n.get("savings_monthly"):
            return None
        d = first_working_day(now.year, now.month)
        hh, _, mm = (n.get("time") or "08:00").partition(":")
        start = datetime(d.year, d.month, d.day, int(hh), int(mm or 0))
        if not (start <= now <= start + timedelta(hours=10)):
            return None
        prev = (d.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
        with _state_lock:
            s = _state()
        if s.get("monthly_done") == prev:
            return None
        tries = s.get("monthly_tries", {})
        if tries.get("month") == prev and (tries.get("count", 0) >= 3 or time.time() - tries.get("at", 0) < 600):
            return None
        return prev

    def run(self, month):
        y, m = int(month[:4]), int(month[5:])
        b = date(y, m, 1).replace(day=28) + timedelta(days=4)
        a, b = date(y, m, 1), b - timedelta(days=b.day)
        ok = True
        for company in self.companies_of():
            try:
                errors = savings.ensure(self.client, company, savings.needed_months(a, b))
                r = savings.report(self.client, company, a, b)
                r["company"] = company
                send(r, self.currency_of(company), actor="scheduler",
                     note=("Some months could not be read: " + "; ".join(errors)) if errors else "")
            except Exception as exc:                               # noqa: BLE001
                traceback.print_exc()
                ok = False
                _log({"ts": datetime.now().isoformat(timespec="seconds"), "actor": "scheduler", "company": company,
                      "status": "failed", "detail": str(exc)[:300], "period": {"from": a.isoformat(),
                                                                            "to": b.isoformat()}})
        with _state_lock:
            s = _state()
            if ok:
                s["monthly_done"] = month
                s.pop("monthly_tries", None)
            else:
                t = s.get("monthly_tries", {})
                count = t.get("count", 0) + 1 if t.get("month") == month else 1
                s["monthly_tries"] = {"month": month, "count": count, "at": time.time()}
                if count >= 3:
                    s["monthly_done"] = month
            _save_state(s)

    def loop(self):
        while True:
            try:
                month = self.due(datetime.now())
                if month:
                    self.run(month)
            except Exception:                                      # noqa: BLE001
                traceback.print_exc()
            time.sleep(60)

    def start(self):
        threading.Thread(target=self.loop, daemon=True, name="savings-monthly").start()
