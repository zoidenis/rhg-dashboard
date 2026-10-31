"""The weekly purchasing report: every Monday, last week's exceptions for the purchasing managers.

One email per company, Monday to Sunday of the week before, with an Excel file of every line:
  * what needs action this week: the points below that cross a limit, in plain sentences;
  * open purchase orders with nothing received yet (what has not reached the unit), by age and by
    location, and the orders received but not invoiced;
  * price changes against the week before (dearer and cheaper, by money impact);
  * contract price coverage: the share of spend on items with an agreed price, items paid above
    it and the largest spend with no agreed price;
  * savings and extra costs against the same dates last year, for the week and the month to date;
  * the same item bought dearer at one location than at another;
  * spend by location, supplier concentration and the critical exceptions nobody owns yet.
Everything is read from Business Central through the same builders as the screens; nothing is
written there. Recipients: Purchasing Manager and Coordinator (To), Cost Controller (Cc) and any
extra address in Settings → Email notifications.
"""
import html
import io
import json
import threading
import time
import traceback
import uuid
from datetime import date, datetime, timedelta

import notify
import operations
import savings
import smtp_config
import store

STATE_FILE = store.DATA_DIR / "weekly_report.json"
_state_lock = threading.Lock()
_jobs, _jobs_lock = {}, threading.Lock()
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _f(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _d(v):
    try:
        d = date.fromisoformat(str(v or "")[:10])
        return None if d.year < 1900 else d
    except ValueError:
        return None


def _m(v):
    return "–" if v is None else f"{v:,.0f}"


def _pct(v):
    return "–" if v is None else f"{v:+.0f}%"


def last_week(today=None):
    today = today or date.today()
    monday = today - timedelta(days=today.weekday())
    return monday - timedelta(days=7), monday - timedelta(days=1)


# ---------------------------------------------------------------- the figures
def build(company, read_dashboard, client, currency, today=None):
    """read_dashboard(company, week_start_iso) -> the Purchasing dashboard payload of that week."""
    today = today or date.today()
    a, b = last_week(today)
    s = store.get_settings()
    th = s["thresholds"]
    errors = []
    data = read_dashboard(company, a.isoformat())
    k = data.get("kpis") or {}
    ops = data.get("operations") or {}
    meta = data.get("meta") or {}

    # ---- open purchase orders, all of them (the screen carries sixty)
    po_rows = []
    try:
        po_rows = client.purchase_order_list(company)
    except Exception as exc:                                       # noqa: BLE001
        errors.append(f"Purchase orders: {str(exc)[:200]}")
    pos = operations.purchase_orders(po_rows, s, today)
    raw = {r.get("No"): r for r in po_rows}
    not_received, to_invoice = [], []
    for p in pos:
        r = raw.get(p["no"]) or {}
        ordered = _d(r.get("Document_Date")) or _d(r.get("Posting_Date"))
        expected = _d(r.get("Requested_Receipt_Date"))
        row = {"no": p["no"], "supplier": p["supplier"] or "", "location": p["location"] or "(no location)",
               "status": p["status"] or "", "ordered": ordered.isoformat() if ordered else "",
               "days": (today - ordered).days if ordered else None,
               "expected": expected.isoformat() if expected else "",
               "late": bool(expected and expected < today),
               "amount": p["amount"], "received_not_invoiced": p["received_not_invoiced"],
               "currency": p["currency"]}
        (to_invoice if p["received_not_invoiced"] > 0 else not_received).append(row)
    not_received.sort(key=lambda r: (-(r["days"] or 0), -r["amount"]))
    to_invoice.sort(key=lambda r: -r["received_not_invoiced"])
    by_loc = {}
    for r in not_received:
        g = by_loc.setdefault(r["location"], {"location": r["location"], "orders": 0, "value": 0.0, "oldest": 0})
        g["orders"] += 1
        g["value"] += r["amount"]
        g["oldest"] = max(g["oldest"], r["days"] or 0)
    po_old_days = int(th.get("po_open_days") or 7)

    # ---- price changes against the week before
    movers = data.get("movers") or []
    ups = sorted([m for m in movers if m["pct"] > 0], key=lambda m: -m["impact"])
    downs = sorted([m for m in movers if m["pct"] < 0], key=lambda m: m["impact"])

    # ---- contract price coverage
    ct = data.get("contracts") or {}
    paid_above = (ops.get("contract_prices") or [])

    # ---- savings and extra costs against last year: the week and the month to date
    sv = {}
    m_from = b.replace(day=1)
    try:
        errs = savings.ensure(client, company, savings.needed_months(min(a, m_from), b))
        if errs:
            errors.append("Savings: " + "; ".join(errs)[:200])
        for key, (x, y) in (("week", (a, b)), ("month", (m_from, b))):
            r = savings.report(client, company, x, y)
            both = [i for i in r["items"] if i["status"] == "both" and i["effect"] is not None]
            sv[key] = {"from": x.isoformat(), "to": y.isoformat(), "summary": r["summary"],
                       "extra": sorted([i for i in both if i["effect"] > 0], key=lambda i: -i["effect"])[:200],
                       "save": sorted([i for i in both if i["effect"] < 0], key=lambda i: i["effect"])[:200],
                       "vendors": sorted(r["vendors"], key=lambda g: -g["net"])[:40]}
    except Exception as exc:                                       # noqa: BLE001
        traceback.print_exc()
        errors.append(f"Savings could not be read: {str(exc)[:200]}")

    # ---- the rest
    spreads = data.get("spreads") or []
    bl = data.get("by_location") or {}
    locations = sorted(({"location": loc, "cur": v.get("cur", 0), "prev": v.get("prev", 0), "pct": v.get("pct"),
                         "documents": v.get("documents", 0), "price_up": (v.get("price_up") or [0, 0])[1]}
                        for loc, v in bl.items() if v.get("cur") or v.get("prev")), key=lambda x: -x["cur"])
    suppliers = ops.get("suppliers") or []
    sla = ops.get("sla") or []
    untracked = data.get("untracked") or {}
    crit = untracked.get("critical") or [e for e in sla if e.get("severity") == "critical"]
    conc = float(th.get("supplier_concentration_pct") or 40)

    # ---- what needs action, in sentences
    acts = []
    old = [r for r in not_received if (r["days"] or 0) > po_old_days]
    if old:
        acts.append(("crit" if len(old) >= 5 else "warn",
                     f"{len(old)} purchase order{'s' if len(old) > 1 else ''} raised more than {po_old_days} days ago with "
                     f"nothing received yet ({_m(sum(r['amount'] for r in old))} {currency}): confirm delivery with "
                     f"the supplier or close the order."))
    late = [r for r in not_received if r["late"]]
    if late:
        acts.append(("warn", f"{len(late)} order{'s' if len(late) > 1 else ''} past the requested receipt date."))
    if to_invoice:
        acts.append(("warn", f"{len(to_invoice)} order{'s' if len(to_invoice) > 1 else ''} received but not invoiced "
                             f"({_m(sum(r['received_not_invoiced'] for r in to_invoice))} {currency}): chase the "
                             f"supplier invoices."))
    if ups:
        acts.append(("crit" if k.get("increase_impact", 0) >= float(th.get("price_critical_impact") or 1e12) else "warn",
                     f"{len(ups)} item{'s' if len(ups) > 1 else ''} bought dearer than the week before: "
                     f"+{_m(k.get('increase_impact'))} {currency}; the largest, {ups[0]['desc'] or ups[0]['item_no']}, "
                     f"{ups[0]['pct']:+.1f}%."))
    if ct.get("available"):
        kp = ct.get("kpis") or {}
        if kp.get("above_value", 0) > 0:
            acts.append(("warn", f"Paid above the agreed price: +{_m(kp['above_value'])} {currency} on the covered spend."))
        if kp.get("coverage_pct", 100) < 70 and ct.get("biggest_uncovered"):
            u = ct["biggest_uncovered"][0]
            acts.append(("info", f"Only {kp['coverage_pct']:.0f}% of the spend has an agreed price; the largest spend "
                                 f"without one is {u['description'] or u['item']} ({_m(u['cost'])} {currency})."))
    if sv.get("month") and sv["month"]["summary"].get("net", 0) > 0:
        acts.append(("warn", f"Month to date, prices cost {_m(sv['month']['summary']['net'])} {currency} more than the "
                             f"same dates last year."))
    if spreads:
        over = sum(x["overpay"] for x in spreads)
        acts.append(("info", f"{len(spreads)} item{'s' if len(spreads) > 1 else ''} bought dearer at one location "
                             f"than at another: {_m(over)} {currency} extra."))
    top_sup = suppliers[0] if suppliers else None
    if top_sup and top_sup.get("share", 0) >= conc:
        acts.append(("info", f"{top_sup['supplier']} holds {top_sup['share']:.0f}% of last week's supplier spend."))
    if crit:
        acts.append(("crit", f"{(untracked.get('counts') or {}).get('critical', len(crit))} critical exception(s) "
                             f"with no owner in the Action center."))

    return {
        "company": company, "currency": currency, "week": {"from": a.isoformat(), "to": b.isoformat()},
        "prior": {"from": (meta.get("prior") or [None, None])[0], "to": (meta.get("prior") or [None, None])[1]},
        "built_at": datetime.now().isoformat(timespec="seconds"), "errors": errors,
        "kpis": k, "actions": acts,
        "not_received": not_received, "not_received_by_location": sorted(by_loc.values(), key=lambda g: -g["value"]),
        "to_invoice": to_invoice, "po_old_days": po_old_days,
        "ups": ups, "downs": downs,
        "contracts": {"available": ct.get("available", False), "kpis": ct.get("kpis") or {},
                      "reason": ct.get("reason", ""), "gaps": ct.get("biggest_gaps") or [],
                      "uncovered": ct.get("biggest_uncovered") or [], "paid_above": paid_above},
        "savings": sv, "spreads": spreads, "locations": locations,
        "suppliers": suppliers[:15], "concentration": conc,
        "exceptions": sorted(sla, key=lambda e: (e.get("severity") != "critical", -abs(e.get("impact") or 0))),
        "untracked_critical": crit[:15],
    }


# ---------------------------------------------------------------- the email
def _period(r):
    a, b = date.fromisoformat(r["week"]["from"]), date.fromisoformat(r["week"]["to"])
    return f"{a:%d %b} – {b:%d %b %Y}"


def subject(r, test=False):
    n = len(r["not_received"])
    return (("[Test] " if test else "") + f"[RHG Purchasing] {r['company']} · Weekly report {_period(r)} · "
            f"{n} open order{'s' if n != 1 else ''} not received")


def render_email(r):
    e = html.escape
    C = r["currency"]
    navy, primary, ink, muted, line, sky = "#08303f", "#0c475e", "#0e2733", "#5c7482", "#d9dcdd", "#e4f3f9"
    green, red, amber = "#0e8a57", "#c0392b", "#b26a00"
    public = (smtp_config.current().get("public_url") or "").rstrip("/")
    th = f"text-align:left;padding:7px 9px;background:{sky};color:{ink};font-size:12px;font-weight:600"
    thn = th.replace("text-align:left", "text-align:right")
    td = f"padding:6px 9px;border-bottom:1px solid #f0f1f2;font-size:12px;color:{ink};vertical-align:top"
    tdn = td + ";text-align:right;white-space:nowrap"
    k = r["kpis"]

    def kpi(label, value, sub, color=ink):
        return (f"<td style='padding:12px 14px;background:#fff;border:1px solid {line};border-radius:10px;"
                f"vertical-align:top;width:25%'><div style='font-size:12px;color:{muted}'>{e(label)}</div>"
                f"<div style='font-size:20px;font-weight:700;color:{color};margin:4px 0'>{e(value)}</div>"
                f"<div style='font-size:11px;color:{muted}'>{e(sub)}</div></td>")

    def h3(t, sub=""):
        return (f"<h3 style='font-size:15px;color:{ink};margin:26px 0 4px'>{e(t)}</h3>"
                + (f"<div style='font-size:12px;color:{muted};margin-bottom:8px'>{e(sub)}</div>" if sub else ""))

    def table(head, rows, empty):
        if not rows:
            return f"<p style='font-size:12px;color:{muted};margin:6px 0'>{e(empty)}</p>"
        hs = "".join(f"<th style='{thn if n else th}'>{e(h)}</th>" for h, n in head)
        return (f"<table role='presentation' cellspacing='0' cellpadding='0' style='width:100%;border-collapse:collapse;"
                f"border:1px solid {line}'><tr>{hs}</tr>" + "".join(rows) + "</table>")

    def link(hash_, text):
        return (f" <a href='{e(public)}/purchasing/#{hash_}' style='color:{primary};font-size:12px'>{e(text)} ›</a>"
                if public else "")
    pct = k.get("total_pct")
    ct = r["contracts"]
    ckp = ct.get("kpis") or {}
    nr = r["not_received"]
    kpis = ("<table role='presentation' cellspacing='8' cellpadding='0' style='width:100%;margin:0 -8px'><tr>"
            + kpi("Spend last week", _m(k.get("total_cur")), (f"{pct:+.1f}% vs the week before" if pct is not None else "")
                  + f" · {C}")
            + kpi("Price changes", f"{k.get('net_price_impact', 0):+,.0f}",
                  f"+{_m(k.get('increase_impact'))} dearer · −{_m(k.get('saving'))} cheaper",
                  red if (k.get("net_price_impact") or 0) > 0 else green)
            + kpi("Contract price coverage", f"{ckp['coverage_pct']:.0f}%" if ct["available"] else "–",
                  f"paid above agreed +{_m(ckp.get('above_value'))}" if ct["available"] else (ct.get("reason") or "")[:60])
            + kpi("Open orders not received", f"{len(nr)}", f"{_m(sum(x['amount'] for x in nr))} {C} ordered",
                  amber if nr else ink)
            + "</tr></table>")
    dot = {"crit": red, "warn": amber, "info": primary}
    actions = ("".join(f"<tr><td style='padding:5px 0;vertical-align:top;width:16px'><span style='display:inline-block;width:8px;"
                       f"height:8px;border-radius:50%;background:{dot[s]};margin-top:5px'></span></td>"
                       f"<td style='padding:4px 0;font-size:13px;color:{ink}'>{e(t)}</td></tr>" for s, t in r["actions"])
               or f"<tr><td style='font-size:13px;color:{muted}'>Nothing crossed a limit last week.</td></tr>")
    actions = f"<table role='presentation' cellspacing='0' cellpadding='0' style='width:100%'>{actions}</table>"

    po_rows = [f"<tr><td style='{td}'><b>{e(x['no'])}</b><div style='color:{muted}'>{e(x['status'])}</div></td>"
               f"<td style='{td}'>{e(x['location'])}</td><td style='{td}'>{e(x['supplier'])}</td>"
               f"<td style='{td}'>{e(x['ordered'])}{(' · due ' + e(x['expected'])) if x['expected'] else ''}</td>"
               f"<td style='{tdn};color:{red if (x['days'] or 0) > r['po_old_days'] else ink}'>{x['days'] if x['days'] is not None else '–'}</td>"
               f"<td style='{tdn}'>{_m(x['amount'])}</td></tr>" for x in nr[:25]]
    loc_rows = [f"<tr><td style='{td}'>{e(g['location'])}</td><td style='{tdn}'>{g['orders']}</td>"
                f"<td style='{tdn}'>{_m(g['value'])}</td><td style='{tdn}'>{g['oldest']}</td></tr>"
                for g in r["not_received_by_location"][:12]]
    inv_rows = [f"<tr><td style='{td}'><b>{e(x['no'])}</b></td><td style='{td}'>{e(x['location'])}</td>"
                f"<td style='{td}'>{e(x['supplier'])}</td><td style='{tdn}'>{x['days'] if x['days'] is not None else '–'}</td>"
                f"<td style='{tdn};color:{amber}'>{_m(x['received_not_invoiced'])}</td></tr>" for x in r["to_invoice"][:12]]

    def mover(m, up):
        return (f"<tr><td style='{td}'>{e(m['desc'] or m['item_no'])} <span style='color:{muted}'>{e(m['item_no'])}</span>"
                f"<div style='color:{muted}'>{e(', '.join(m.get('locations') or [])[:80])}</div></td>"
                f"<td style='{tdn}'>{m['prev_avg']:,.2f} → {m['cur_avg']:,.2f}</td>"
                f"<td style='{tdn};color:{red if up else green}'>{m['pct']:+.1f}%</td>"
                f"<td style='{tdn}'><b>{m['impact']:+,.0f}</b></td></tr>")
    mv_head = [("Item", 0), ("Unit cost", 1), ("Change", 1), ("Impact", 1)]

    gap_rows = [f"<tr><td style='{td}'>{e(x.get('desc') or x.get('item_no'))}<div style='color:{muted}'>{e(x.get('vendor') or '')}</div></td>"
                f"<td style='{tdn}'>{_f(x.get('contract')):,.2f}</td><td style='{tdn}'>{_f(x.get('paid')):,.2f}</td>"
                f"<td style='{tdn};color:{red}'>{_f(x.get('pct')):+.1f}%</td><td style='{tdn}'><b>{_m(x.get('overpaid'))}</b></td></tr>"
                for x in ct["paid_above"][:10]]
    unc_rows = [f"<tr><td style='{td}'>{e(x['description'] or x['item'])}<div style='color:{muted}'>{e(x.get('vendor') or '')}</div></td>"
                f"<td style='{td}'>{e(x.get('category') or '')}</td><td style='{tdn}'>{_m(x['cost'])}</td></tr>"
                for x in ct["uncovered"][:8]]

    def sv_block(key, title):
        x = r["savings"].get(key)
        if not x:
            return ""
        sm = x["summary"]
        rows = ([f"<tr><td style='{td}'>{e(i['desc'] or i['item'])}<div style='color:{muted}'>{e(', '.join(i['vendors'][:2]))}</div></td>"
                 f"<td style='{tdn}'>{(i['ly_avg'] or 0):,.2f} → {(i['cur_avg'] or 0):,.2f}</td>"
                 f"<td style='{tdn};color:{red}'>{i['pct']:+.1f}%</td><td style='{tdn}'><b>+{_m(i['effect'])}</b></td></tr>"
                 for i in x["extra"][:6]]
                + [f"<tr><td style='{td}'>{e(i['desc'] or i['item'])}<div style='color:{muted}'>{e(', '.join(i['vendors'][:2]))}</div></td>"
                   f"<td style='{tdn}'>{(i['ly_avg'] or 0):,.2f} → {(i['cur_avg'] or 0):,.2f}</td>"
                   f"<td style='{tdn};color:{green}'>{i['pct']:+.1f}%</td><td style='{tdn}'><b>{_m(i['effect'])}</b></td></tr>"
                   for i in x["save"][:4]])
        net = sm.get("net", 0)
        return (f"<div style='font-size:13px;margin:8px 0 6px'><b>{e(title)}</b> ({e(x['from'])} → {e(x['to'])}): "
                f"savings <b style='color:{green}'>{_m(sm.get('save'))}</b> · extra <b style='color:{red}'>{_m(sm.get('extra'))}</b> · "
                f"net <b style='color:{red if net > 0 else green}'>{'+' if net > 0 else '−'}{_m(abs(net))}</b> {e(C)} "
                f"<span style='color:{muted}'>(spend {_m(sm.get('ly'))} last year → {_m(sm.get('cur'))})</span></div>"
                + table([("Item", 0), ("Unit cost, last year → now", 1), ("Change", 1), ("Effect", 1)], rows,
                        "No item bought in both periods."))
    spread_rows = [f"<tr><td style='{td}'>{e(x['desc'] or x['item_no'])}</td><td style='{td}'>{e(x['low_loc'])} {x['low']:,.2f}</td>"
                   f"<td style='{td}'>{e(x['high_loc'])} {x['high']:,.2f}</td><td style='{tdn}'>{x['spread']:+.0f}%</td>"
                   f"<td style='{tdn};color:{red}'>{_m(x['overpay'])}</td></tr>" for x in r["spreads"][:8]]
    loc_spend = [f"<tr><td style='{td}'>{e(x['location'])}</td><td style='{tdn}'>{_m(x['cur'])}</td><td style='{tdn}'>{_m(x['prev'])}</td>"
                 f"<td style='{tdn};color:{red if (x['pct'] or 0) > 15 else ink}'>{_pct(x['pct'])}</td>"
                 f"<td style='{tdn}'>{('+' + _m(x['price_up'])) if x['price_up'] else '–'}</td></tr>" for x in r["locations"][:14]]
    sup_rows = [f"<tr><td style='{td}'>{e(x['supplier'])}</td><td style='{tdn}'>{_m(x['spend'])}</td>"
                f"<td style='{tdn};color:{red if x['share'] >= r['concentration'] else ink}'>{x['share']:.0f}%</td>"
                f"<td style='{tdn}'>{x['invoices']}</td></tr>" for x in r["suppliers"][:8]]
    exc_rows = [f"<tr><td style='{td}'><span style='color:{red}'>●</span> <b>{e(x.get('title') or '')}</b>"
                f"<div style='color:{muted}'>{e(x.get('detail') or '')}</div></td><td style='{td}'>{e(x.get('location') or '')}</td>"
                f"<td style='{tdn}'>{_m(abs(x.get('impact') or 0))}</td></tr>" for x in r["untracked_critical"][:10]]
    errs = "".join(f"<div style='background:#fdf3e2;color:{amber};padding:8px 12px;border-radius:8px;font-size:12px;"
                   f"margin-bottom:10px'>{e(x)}</div>" for x in r["errors"])
    button = (f"<p style='margin:24px 0 0'><a href='{e(public)}/purchasing/' style='background:{primary};color:#fff;"
              f"text-decoration:none;padding:10px 18px;border-radius:8px;font-size:13px;font-weight:600;display:inline-block'>"
              f"Open the Purchasing tower</a></p>") if public else ""
    body = f"""<!doctype html><html><head><meta charset="utf-8"></head><body style="margin:0;background:#ecebeb;font-family:Figtree,Segoe UI,Arial,sans-serif;color:{ink}">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="background:#ecebeb"><tr><td align="center" style="padding:24px 12px">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="max-width:900px">
<tr><td style="background:{navy};color:#fff;padding:18px 22px;border-radius:12px 12px 0 0">
<div style="font-size:11px;letter-spacing:.12em;color:#a0d8ea">REVO HOSPITALITY GROUP · RHG PURCHASING</div>
<div style="font-size:20px;font-weight:700;margin-top:4px">Weekly purchasing report</div>
<div style="font-size:12px;color:#cfe3ea;margin-top:4px">{e(r['company'])} · {e(_period(r))} · against {e(r['prior']['from'] or '')} → {e(r['prior']['to'] or '')} · open orders as of {e(r['built_at'][:10])}</div></td></tr>
<tr><td style="background:#fff;padding:22px;border-radius:0 0 12px 12px">
{errs}{kpis}
{h3("What needs action this week")}{actions}
{h3(f"Open purchase orders, nothing received yet ({len(nr)})", f"Orders still open in Business Central with no goods received: what has not reached the unit. Oldest first; days since the order date, red beyond {r['po_old_days']} days.")}
{table([("Order", 0), ("Location", 0), ("Supplier", 0), ("Ordered", 0), ("Days", 1), (f"Amount {C}", 1)], po_rows, "No open order waiting for goods.")}
{('<div style="font-size:11px;color:' + muted + ';margin-top:4px">' + str(len(nr) - 25) + ' more in the Excel file.</div>') if len(nr) > 25 else ''}
{h3("By location") if loc_rows else ''}{table([("Location", 0), ("Orders", 1), (f"Amount {C}", 1), ("Oldest, days", 1)], loc_rows, "") if loc_rows else ''}
{h3(f"Received, not invoiced ({len(r['to_invoice'])})", "Goods are in, the supplier invoice is not posted: the cost is accrued but not confirmed.")}
{table([("Order", 0), ("Location", 0), ("Supplier", 0), ("Days open", 1), (f"Not invoiced {C}", 1)], inv_rows, "Nothing received and waiting for an invoice.")}
{h3("Price changes against the week before", "Average unit cost of items bought in both weeks; impact = change × last week's quantity.")}{link("prices", "All price changes")}
{table(mv_head, [mover(m, True) for m in r["ups"][:10]], "No item bought dearer than the week before.")}
<div style="height:8px"></div>{table(mv_head, [mover(m, False) for m in r["downs"][:5]], "No item bought cheaper than the week before.")}
{h3("Contract price coverage", (f"{ckp.get('coverage_pct', 0):.0f}% of last week's spend is on items with an agreed price ({_m(ckp.get('covered_cost'))} of {_m(ckp.get('total_cost'))} {C}); {ckp.get('items_covered', 0)} of {ckp.get('items_bought', 0)} items. Paid above agreed: +{_m(ckp.get('above_value'))}; below: {_m(ckp.get('below_value'))}.") if ct['available'] else (ct.get('reason') or 'The price list could not be read.'))}{link("contracts", "Contracts")}
{table([("Item paid above the agreed price", 0), ("Agreed", 1), ("Paid", 1), ("Gap", 1), (f"Overpaid {C}", 1)], gap_rows, "No item paid above its agreed price.") if ct['available'] else ''}
{('<div style="height:8px"></div>' + table([("Largest spend with no agreed price", 0), ("Category", 0), (f"Spend {C}", 1)], unc_rows, "")) if unc_rows else ''}
{h3("Savings and extra costs against last year", "Price effect on items bought in both periods: (unit cost now − same dates last year) × quantity now.")}{link("savings", "Price savings")}
{sv_block("week", "Last week")}{sv_block("month", "Month to date")}
{'' if r['savings'] else '<p style="font-size:12px;color:' + muted + '">Savings could not be read this time.</p>'}
{h3("Same item, dearer at another location", "Extra paid = the dearest location's quantity priced at the cheapest location's cost.") if spread_rows else ''}
{table([("Item", 0), ("Cheapest", 0), ("Dearest", 0), ("Spread", 1), (f"Extra paid {C}", 1)], spread_rows, "") if spread_rows else ''}
{h3("Spend by location")}{table([("Location", 0), ("Last week", 1), ("Week before", 1), ("Change", 1), ("Price rises", 1)], loc_spend, "No purchases last week.")}
{h3("Largest suppliers", f"Share of last week's supplier invoices; red from {r['concentration']:.0f}%.")}{table([("Supplier", 0), (f"Spend {C}", 1), ("Share", 1), ("Invoices", 1)], sup_rows, "No supplier invoice last week.")}
{h3("Critical exceptions with no owner", "Turn them into actions in the Action center so somebody owns them.")}{link("exceptions", "Exceptions")}
{table([("Exception", 0), ("Location", 0), (f"Impact {C}", 1)], exc_rows, "Every critical exception has an owner.")}
<p style="color:{muted};font-size:12px;margin:18px 0 0"><b>The attached Excel file lists every open order, price change, contract gap, saving and exception.</b> Amounts in {e(C)}; purchase cost excludes VAT.</p>
{button}
<p style="color:{muted};font-size:11px;margin:24px 0 0;border-top:1px solid {line};padding-top:12px">Sent every week by RHG Control Towers from Business Central; nothing was changed there. Recipients and the day are in Purchasing → Settings → Email notifications.</p>
</td></tr></table></td></tr></table></body></html>"""
    text = (f"RHG Purchasing · Weekly report · {r['company']} · {_period(r)}\n\n"
            + "\n".join(f"- {t}" for _, t in r["actions"])
            + f"\n\nOpen orders not received: {len(nr)} ({_m(sum(x['amount'] for x in nr))} {C})"
            + f"\nReceived, not invoiced: {len(r['to_invoice'])}"
            + f"\nSpend last week: {_m(k.get('total_cur'))} {C}"
            + "\n\nThe attached Excel file lists every line." + (f"\n\nOpen: {public}/purchasing/" if public else ""))
    return body, text


def render_xlsx(r):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    fill, font = PatternFill("solid", fgColor="E4F3F9"), Font(bold=True, color="0E2733")
    money, unit, pct = '#,##0;[Red]-#,##0', '#,##0.00', '0.0'

    def sheet(title, header, rows, formats, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font, c.alignment = fill, font, Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append(row)
        for i, f in enumerate(formats, start=1):
            if f:
                for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell[0].number_format = f
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws
    lvl = {"crit": "Critical", "warn": "Warning", "info": "Info"}
    sheet("Actions", ["Priority", "What needs action"], [[lvl[s], t] for s, t in r["actions"]], [None, None], [12, 120],
          first=True)
    sheet("Open orders not received", ["Order", "Status", "Location", "Supplier", "Order date", "Days open",
                                       "Requested receipt", "Past due", "Amount", "Currency"],
          [[x["no"], x["status"], x["location"], x["supplier"], x["ordered"], x["days"], x["expected"],
            "yes" if x["late"] else "", round(x["amount"], 2), x["currency"]] for x in r["not_received"]],
          [None] * 8 + [money, None], [12, 10, 16, 32, 12, 10, 14, 9, 14, 8])
    sheet("Received not invoiced", ["Order", "Location", "Supplier", "Days open", "Received, not invoiced", "Order amount"],
          [[x["no"], x["location"], x["supplier"], x["days"], round(x["received_not_invoiced"], 2), round(x["amount"], 2)]
           for x in r["to_invoice"]], [None, None, None, None, money, money], [12, 16, 32, 10, 18, 14])
    sheet("Price changes", ["Item", "Description", "Locations", "Unit cost week before", "Unit cost last week", "Change %",
                            "Quantity", "Impact"],
          [[m["item_no"], m["desc"], ", ".join(m.get("locations") or []), round(m["prev_avg"], 4), round(m["cur_avg"], 4),
            round(m["pct"], 1), round(m.get("cur_qty") or 0, 3), round(m["impact"], 0)] for m in r["ups"] + r["downs"]],
          [None, None, None, unit, unit, pct, unit, money], [12, 34, 26, 14, 14, 10, 12, 14])
    ct = r["contracts"]
    sheet("Paid above contract", ["Item", "Description", "Supplier", "Agreed price", "Paid", "Gap %", "Quantity", "Overpaid"],
          [[x.get("item_no"), x.get("desc"), x.get("vendor"), round(_f(x.get("contract")), 4), round(_f(x.get("paid")), 4),
            round(_f(x.get("pct")), 1), round(_f(x.get("qty")), 3), round(_f(x.get("overpaid")), 0)] for x in ct["paid_above"]],
          [None, None, None, unit, unit, pct, unit, money], [12, 34, 26, 12, 12, 9, 12, 14])
    sheet("No agreed price", ["Item", "Description", "Supplier", "Category", "Spend"],
          [[x["item"], x["description"], x.get("vendor"), x.get("category"), round(x["cost"], 0)] for x in ct["uncovered"]],
          [None, None, None, None, money], [12, 34, 26, 20, 14])
    for key, title in (("week", "Savings last week"), ("month", "Savings month to date")):
        x = r["savings"].get(key)
        if not x:
            continue
        sheet(title, ["Item", "Description", "Suppliers", "Unit cost last year", "Unit cost now", "Change %", "Effect"],
              [[i["item"], i["desc"], ", ".join(i["vendors"][:3]), i["ly_avg"], i["cur_avg"], i["pct"], i["effect"]]
               for i in x["extra"] + x["save"]], [None, None, None, unit, unit, pct, money], [12, 34, 30, 14, 14, 10, 14])
    sheet("Location price spread", ["Item", "Description", "Cheapest location", "Cost", "Dearest location", "Cost",
                                    "Spread %", "Extra paid"],
          [[x["item_no"], x["desc"], x["low_loc"], round(x["low"], 4), x["high_loc"], round(x["high"], 4),
            round(x["spread"], 1), round(x["overpay"], 0)] for x in r["spreads"]],
          [None, None, None, unit, None, unit, pct, money], [12, 34, 16, 12, 16, 12, 10, 14])
    sheet("Spend by location", ["Location", "Last week", "Week before", "Change %", "Documents", "Price rises"],
          [[x["location"], round(x["cur"], 0), round(x["prev"], 0), None if x["pct"] is None else round(x["pct"], 1),
            x["documents"], round(x["price_up"], 0)] for x in r["locations"]],
          [None, money, money, pct, None, money], [18, 14, 14, 10, 11, 14])
    sheet("Exceptions", ["Severity", "Rule", "Exception", "Detail", "Location", "Supplier", "Impact", "Next action"],
          [[x.get("severity"), x.get("rule"), x.get("title"), x.get("detail"), x.get("location"), x.get("supplier"),
            round(abs(x.get("impact") or 0), 0), x.get("next_action")] for x in r["exceptions"]],
          [None] * 6 + [money, None], [10, 26, 40, 60, 18, 24, 14, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- sending, history, schedule
def recipients(settings=None):
    n = (settings or store.get_settings()).get("notifications") or {}
    to = (notify.parse_emails(n.get("manager", ""))[0] + notify.parse_emails(n.get("coordinator", ""))[0]
          + notify.parse_emails(n.get("weekly_to", ""))[0])
    cc = [c for c in notify.parse_emails(n.get("cost_controller", ""))[0] if c not in to]
    return list(dict.fromkeys(to)), cc


def send(r, to=None, cc=None, actor="scheduler", test=False):
    if to is None:
        to, cc = recipients()
    if not to:
        raise RuntimeError("No Purchasing Manager, Coordinator or weekly report address is set in Settings → Email notifications.")
    body, text = render_email(r)
    stem = f"RHG_{r['company'].replace(' ', '_')}_purchasing_week_{r['week']['from']}"
    notify.send_mail(to, cc or [], subject(r, test), body, text,
                     [(f"{stem}.xlsx", "application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       render_xlsx(r))])
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), "actor": actor, "company": r["company"],
             "week": r["week"], "to": to, "cc": cc or [], "status": "sent", "test": test,
             "detail": f"{len(r['not_received'])} orders not received · {len(r['actions'])} points"}
    _log(entry)
    store.audit("email", "weekly purchasing report", actor, {}, entry, subject(r, test))
    return entry


def _state():
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_state(s):
    tmp = STATE_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(s, indent=1, default=str), encoding="utf-8")
    tmp.replace(STATE_FILE)


def _log(entry):
    with _state_lock:
        s = _state()
        s["runs"] = (s.get("runs", []) + [entry])[-100:]
        _save_state(s)


def recent_runs(limit=15):
    with _state_lock:
        return list(reversed(_state().get("runs", [])))[:limit]


def _when(n):
    try:
        day = max(0, min(6, int(float(n.get("weekly_day") or 0))))
    except (TypeError, ValueError):
        day = 0
    hh, _, mm = (n.get("time") or "08:00").partition(":")
    return day, int(hh), int(mm or 0)


def next_run(now=None):
    n = store.get_settings().get("notifications") or {}
    if not n.get("weekly_report"):
        return None
    now = now or datetime.now()
    day, hh, mm = _when(n)
    for k in range(8):
        d = now.date() + timedelta(days=k)
        if d.weekday() != day:
            continue
        t = datetime(d.year, d.month, d.day, hh, mm)
        week = last_week(d)[0].isoformat()
        with _state_lock:
            done = _state().get("done")
        if t + timedelta(hours=10) >= now and done != week:
            return t.isoformat(timespec="minutes")
    return None


class Weekly:
    """Sends last week's report on the day and at the time in Settings (Monday by default)."""

    def __init__(self, read_dashboard, client, companies_of, currency_of):
        self.read, self.client, self.companies_of, self.currency_of = read_dashboard, client, companies_of, currency_of

    def due(self, now):
        n = store.get_settings().get("notifications") or {}
        if not n.get("weekly_report"):
            return None
        day, hh, mm = _when(n)
        if now.weekday() != day:
            return None
        start = datetime(now.year, now.month, now.day, hh, mm)
        if not (start <= now <= start + timedelta(hours=10)):
            return None
        week = last_week(now.date())[0].isoformat()
        with _state_lock:
            s = _state()
        if s.get("done") == week:
            return None
        tries = s.get("tries", {})
        if tries.get("week") == week and (tries.get("count", 0) >= 3 or time.time() - tries.get("at", 0) < 900):
            return None
        return week

    def run(self, week):
        ok = True
        for company in self.companies_of():
            try:
                send(build(company, self.read, self.client, self.currency_of(company)))
            except Exception as exc:                                # noqa: BLE001
                traceback.print_exc()
                ok = False
                _log({"ts": datetime.now().isoformat(timespec="seconds"), "actor": "scheduler", "company": company,
                      "status": "failed", "detail": str(exc)[:300], "week": {"from": week}})
        with _state_lock:
            s = _state()
            if ok:
                s["done"] = week
                s.pop("tries", None)
            else:
                t = s.get("tries", {})
                count = t.get("count", 0) + 1 if t.get("week") == week else 1
                s["tries"] = {"week": week, "count": count, "at": time.time()}
                if count >= 3:
                    s["done"] = week
            _save_state(s)

    def loop(self):
        while True:
            try:
                week = self.due(datetime.now())
                if week:
                    self.run(week)
            except Exception:                                      # noqa: BLE001
                traceback.print_exc()
            time.sleep(60)

    def start(self):
        threading.Thread(target=self.loop, daemon=True, name="weekly-report").start()

    # ---- preview and test, in the background: a week of BC reads takes a while
    def start_job(self, company, actor, send_to=None):
        with _jobs_lock:
            for j in _jobs.values():
                if j["actor"] == actor and j["status"] == "running":
                    return dict(j, html=None)
            if len(_jobs) > 20:
                for key in sorted(_jobs, key=lambda x: _jobs[x]["started"])[:10]:
                    if _jobs[key]["status"] != "running":
                        _jobs.pop(key, None)
            job = {"id": uuid.uuid4().hex, "actor": actor, "company": company, "to": send_to or [],
                   "status": "running", "started": time.time(), "detail": "", "html": None}
            _jobs[job["id"]] = job

        def go():
            try:
                r = build(company, self.read, self.client, self.currency_of(company))
                job["html"] = render_email(r)[0]
                if send_to:
                    send(r, to=send_to, cc=[], actor=actor, test=True)
                    job["detail"] = "Sent to " + ", ".join(send_to)
                job["status"] = "done"
            except Exception as exc:                               # noqa: BLE001
                traceback.print_exc()
                job.update(status="failed", detail=str(exc)[:300])
                if send_to:
                    _log({"ts": datetime.now().isoformat(timespec="seconds"), "actor": actor, "company": company,
                          "status": "failed", "detail": str(exc)[:300], "test": True})
        threading.Thread(target=go, daemon=True, name="weekly-report-job").start()
        return dict(job, html=None)

    @staticmethod
    def job(job_id, actor):
        with _jobs_lock:
            j = _jobs.get(job_id)
            return j if j and j["actor"] == actor else None
