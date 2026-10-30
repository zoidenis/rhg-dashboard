"""Fiscal sales against Business Central: the e-invoices we issued, as the Tax Administration
lists them (e-invoice service, PartyType SELLER), next to the sales invoices and credit memos
posted in BC (published as PBSI1 / PBSC1, one row per line, with the NIVF and the customer's NIPT).

Scope: e-invoices, i.e. invoices to businesses and public bodies. Sales to individuals at the
till (POS, customer "INDIVID") are fiscalized by the till as cash invoices; the e-invoice service
does not list them, so they are not part of this check (Sales · Simphony covers the takings).

Matching, each BC document used once:
1. EIC = NIVF on a posted sales invoice (positive) or credit memo (negative); amounts compared in
   the e-invoice's currency (lek: BC's lek amount; EUR: the BC document in EUR).
2. Buyer NIPT + amount on a posted document without NIVF, posting dates within 45 days.
What is left is missing in BC (an e-invoice issued outside BC, e.g. from the Self-Care portal),
or recent while younger than the grace days. An e-invoice the buyer REFUSED is flagged: if it is
still in BC, it needs a credit memo.

From the BC side: invoices to customers with an Albanian NIPT (B2B) without an e-invoice code,
BC codes the Tax Administration does not list for us, and the same code on two documents.

Business Central is only read; only reading requests go to the Tax Administration.
"""
import io
import re
import threading
import time
import traceback
from datetime import date, datetime, timedelta

import cis_client
import config
import fiscal
import store

NUIS = re.compile(r"^[A-Z]\d{8}[A-Z]$")
SIDE = "sales"


def _key(company):
    return f"{SIDE}:{company}"


def coverage(company):
    with fiscal._lock:
        db = fiscal._db()
        cov = db.execute("SELECT source, date_from, date_to, rows, at FROM coverage WHERE company=? AND "
                         "source LIKE '%:sales' ORDER BY at DESC", (company,)).fetchall()
        n = db.execute("SELECT COUNT(*) FROM einv_sales WHERE company=?", (company,)).fetchone()[0]
    spans = sorted((c[1], c[2]) for c in cov if c[1] and c[2])
    merged = []
    for a, b in spans:
        if merged and a <= (date.fromisoformat(merged[-1][1]) + timedelta(days=1)).isoformat():
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    last = merged[-1] if merged else [None, None]
    return {"loads": [dict(zip(("source", "from", "to", "rows", "at"), (c[0].split(":")[0],) + tuple(c[1:])))
                      for c in cov[:12]], "spans": merged, "from": last[0], "to": last[1], "count": n}


def _store(company, rows, span):
    now = datetime.now().isoformat(timespec="seconds")
    with fiscal._lock:
        db = fiscal._db()
        db.executemany("INSERT OR REPLACE INTO einv_sales VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                       [(company, r["eic"].lower(), r.get("inv_no"), r.get("date"),
                         (r.get("buyer_nipt") or "").strip().upper(), r.get("buyer_name") or "", r.get("total"),
                         (r.get("currency") or "ALL").upper(), r.get("status"), (r.get("due") or "")[:10],
                         "cis", now) for r in rows if r.get("eic")])
        db.execute("INSERT INTO coverage VALUES (?,?,?,?,?,?,?)", (company, "cis:sales", span[0], span[1], len(rows),
                                                                  now, ""))
        db.commit()
    with fiscal._jobs_lock:
        for k in [k for k in fiscal._results if k[0] == _key(company)]:
            fiscal._results.pop(k, None)


def fetch(company, days=None, actor="scheduler", full=False):
    """Reads the e-invoices we issued (PartyType SELLER) from the e-invoice service."""
    days = int(days or fiscal._num(fiscal.settings(), "days", 45))
    b = date.today()
    a = b - timedelta(days=days)
    with fiscal._state_lock:
        last = (fiscal._read_state().get("fetched_to") or {}).get(_key(company))
    if last and not full:
        a = max(a, date.fromisoformat(last) - timedelta(days=2))
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), "kind": "fetch sales", "company": company,
             "actor": actor, "from": a.isoformat(), "to": b.isoformat()}
    try:
        rows = cis_client.get_einvoices(a, b, party="SELLER")
        own = (config.get("FISK_NIPT", "") or "").strip().upper()
        if own:
            rows = [r for r in rows if not r.get("seller_nipt") or r["seller_nipt"].upper() == own]
        fiscal._names(rows, nipt_key="buyer_nipt", name_key="buyer_name")
        _store(company, rows, (a.isoformat(), b.isoformat()))
        with fiscal._state_lock:
            st = fiscal._read_state()
            st.setdefault("fetched_to", {})[_key(company)] = b.isoformat()
            fiscal._write_state(st)
        entry.update(status="ok", rows=len(rows))
    except cis_client.NotConfigured as exc:
        entry.update(status="not configured", detail=str(exc))
    except Exception as exc:                                     # noqa: BLE001
        traceback.print_exc()
        entry.update(status="failed", detail=str(exc)[:400])
    fiscal._log(entry)
    store.audit("fiscal fetch sales", company, actor, {},
                {k: entry.get(k) for k in ("status", "rows", "from", "to")}, entry.get("detail", ""))
    return entry


# ---------------------------------------------------------------- Business Central
def _customer_nipts(client, company, numbers):
    """NIPT per customer, from the sales lines (the customer list is not published in every
    company). Kept for a week in the fiscal state."""
    with fiscal._state_lock:
        cache = (fiscal._read_state().get("customer_nipt") or {}).get(company, {})
    now = time.time()
    out, ent = {}, config.get("E_SALES_INV", "PBSI1")
    for no in numbers:
        hit = cache.get(no)
        if hit and now - hit[1] < 7 * 86400:
            out[no] = hit[0]
            continue
        nipt = ""
        for e in (ent, config.get("E_SALES_CM", "PBSC1")):
            try:
                rows = client._get_all(client._url(company, e), {"$filter": f"Sell_to_Customer_No_ eq '{no.replace(chr(39), chr(39) * 2)}'",
                                                                 "$select": "Sell_to_Customer_No_,VAT_Registration_No_",
                                                                 "$top": 1}, limit=1)
            except Exception as exc:                             # noqa: BLE001
                print(f"fiscal sales: NIPT of {no} not read: {str(exc)[:120]}")
                rows = []
            if rows:
                nipt = (rows[0].get("VAT_Registration_No_") or "").strip().upper()
                break
        out[no] = nipt
        cache[no] = [nipt, now]
    with fiscal._state_lock:
        st = fiscal._read_state()
        st.setdefault("customer_nipt", {})[company] = cache
        fiscal._write_state(st)
    return out


def _bc(client, company, a):
    """Posted sales documents (customer ledger: invoices and credit memos with document and lek
    amounts), their e-invoice codes (sales lines carrying a NIVF) and the customers' NIPTs."""
    start = (date.fromisoformat(a) - timedelta(days=60)).isoformat()
    cle = client._get_all(client._url(company, config.E_CUSTOMER), {
        "$filter": f"Posting_Date ge {start} and (Document_Type eq 'Invoice' or Document_Type eq 'Credit Memo')",
        "$select": "Document_No,Document_Type,Customer_No,Customer_Name,Posting_Date,Due_Date,Currency_Code,Amount,"
                   "Amount_LCY"})
    eic = {}
    for e, kind in ((config.get("E_SALES_INV", "PBSI1"), "invoice"), (config.get("E_SALES_CM", "PBSC1"), "credit memo")):
        for r in client._get_all(client._url(company, e), {
                "$filter": f"Posting_Date ge {start} and NIVF ne ''",
                "$select": "Document_No_,Sell_to_Customer_No_,VAT_Registration_No_,NIVF"}):
            eic.setdefault((kind, r.get("Document_No_")), ((r.get("NIVF") or "").strip().lower(),
                                                          (r.get("VAT_Registration_No_") or "").strip().upper()))
    nipts = _customer_nipts(client, company, sorted({c.get("Customer_No") for c in cle if c.get("Customer_No")}))
    docs = []
    for c in cle:
        kind = "invoice" if c.get("Document_Type") == "Invoice" else "credit memo"
        code, nipt = eic.get((kind, c.get("Document_No")), ("", ""))
        cur = (c.get("Currency_Code") or "").upper()
        docs.append({"no": c.get("Document_No"), "kind": kind, "customer_no": c.get("Customer_No") or "",
                     "customer": c.get("Customer_Name") or "", "nipt": nipt or nipts.get(c.get("Customer_No"), ""),
                     "currency": cur, "posting_date": c.get("Posting_Date"), "due": c.get("Due_Date"), "eic": code,
                     "amount": round(float(c.get("Amount") or 0), 2), "lcy": round(float(c.get("Amount_LCY") or 0), 2),
                     "location": "", "description": ""})
    return docs, len(cle)


def reconcile(einv, docs, a, b, today, s):
    tol = fiscal._num(s, "tolerance", 1)
    grace = int(fiscal._num(s, "grace_days", 7))
    own = (config.get("FISK_NIPT", "") or "").strip().upper()
    by_eic = {}
    for d in docs:
        if d["eic"]:
            by_eic.setdefault(d["eic"], []).append(d)
    by_nipt = {}
    for d in docs:
        for k in fiscal.nipt_keys(d["nipt"]):
            by_nipt.setdefault(k, []).append(d)
    name_of = {}
    for d in docs:
        for k in fiscal.nipt_keys(d["nipt"]):
            name_of.setdefault(k, d["customer"])

    def amount_for(d, cur):
        """BC amount, unsigned (the service lists credit notes with a positive amount too)."""
        if cur == "ALL":
            return abs(d["lcy"])
        return abs(d["amount"]) if d["currency"] == cur else None

    def near(d1, d2):
        try:
            return abs((date.fromisoformat(str(d1)[:10]) - date.fromisoformat(str(d2)[:10])).days)
        except (TypeError, ValueError):
            return 999
    used, rows = set(), []
    for e in einv:
        total = abs(e["total"] or 0.0)
        cur = (e.get("currency") or "ALL").upper()
        cur = "ALL" if cur in ("", "LEK") else cur
        buyer = e["buyer_name"] or next((name_of[k] for k in fiscal.nipt_keys(e["buyer_nipt"]) if k in name_of), "")
        r = {"eic": e["eic"], "date": e["inv_date"], "buyer": buyer, "nipt": e["buyer_nipt"] or "", "inv_no": e["inv_no"] or "",
             "total": round(total, 2), "currency": cur, "status_tax": (e.get("status") or "").upper(), "due": e.get("due") or "",
             "bc_no": "", "bc_customer": "", "bc_amount": None, "bc_date": "", "diff": None, "status": "", "note": "",
             "value": round(total, 2) if cur == "ALL" else 0.0}
        m = [d for d in by_eic.get(e["eic"], []) if d["no"] not in used]
        if not m:
            cand = [d for k in fiscal.nipt_keys(r["nipt"]) for d in by_nipt.get(k, []) if d["no"] not in used
                    and not d["eic"] and amount_for(d, cur) is not None and abs(amount_for(d, cur) - total) <= tol
                    and near(d["posting_date"], r["date"]) <= 45]
            cand.sort(key=lambda d: near(d["posting_date"], r["date"]))
            m = cand[:1]
            via = "nipt"
        else:
            via = "eic"
        if m:
            d = m[0]
            used.add(d["no"])
            amt = amount_for(d, cur)
            r.update(bc_no=d["no"], bc_customer=d["customer"], bc_date=d["posting_date"] or "",
                     bc_amount=round(amt, 2) if amt is not None else None)
            r["value"] = d["lcy"]                              # signed: credit memos count negative
            if d["kind"] == "credit memo":
                r["note"] = "credit memo"
            if via == "nipt":
                r["status"] = "matched, no EIC in BC"
            elif amt is None:
                r.update(status="matched", note=f"e-invoice in {cur}, BC in {d['currency'] or 'ALL'}: amount not compared")
            elif abs(total - amt) <= (max(tol, abs(total) * 0.01) if d["currency"] and cur == "ALL" else tol):
                r["status"] = "matched"
            else:
                r.update(status="amount differs", diff=round(total - amt, 2))
            if r["status_tax"] == "REFUSED":
                r.update(status="refused by buyer", note="refused by the buyer but posted in BC: a credit memo is due")
        elif r["status_tax"] == "REFUSED":
            r.update(status="refused by buyer", note="refused by the buyer, not in BC")
        elif r["date"] and near(r["date"], today) <= grace:
            r["status"] = "recent"
        else:
            r["status"] = "missing"
            r["note"] = "issued at the Tax Administration, not found in BC"
        rows.append(r)

    eics = {e["eic"] for e in einv}
    in_period = [d for d in docs if a <= (d["posting_date"] or "")[:10] <= b]
    dup = []
    for k, ds in by_eic.items():
        if len(ds) > 1:
            same = len({round(x["lcy"], 2) for x in ds}) == 1
            for x in ds:
                dup.append({"eic": k, "bc_no": x["no"], "customer": x["customer"], "posting_date": x["posting_date"],
                            "amount": x["lcy"], "kind": "same amount" if same else "different amounts",
                            "in_fiscal": k in eics})
    dup.sort(key=lambda x: (x["kind"] != "same amount", -abs(x["amount"]), x["eic"]))
    not_fiscal = [{"bc_no": d["no"], "customer": d["customer"], "nipt": d["nipt"], "eic": d["eic"],
                   "posting_date": d["posting_date"], "amount": d["lcy"], "kind": d["kind"]}
                  for d in in_period if d["eic"] and d["eic"] not in eics]
    b2b = [{"bc_no": d["no"], "customer": d["customer"], "customer_no": d["customer_no"], "nipt": d["nipt"],
            "posting_date": d["posting_date"], "amount": d["lcy"], "currency": d["currency"], "kind": d["kind"],
            "location": d["location"], "description": d["description"]}
           for d in in_period if not d["eic"] and d["no"] not in used and NUIS.match(d["nipt"] or "")
           and d["nipt"] != own]
    b2b.sort(key=lambda x: -abs(x["amount"]))
    not_fiscal.sort(key=lambda x: -abs(x["amount"]))

    def tally(st):
        sel = [r for r in rows if r["status"] == st]
        return [len(sel), round(sum(r["value"] for r in sel), 0)]
    k = {st: tally(st) for st in ("matched", "matched, no EIC in BC", "amount differs", "refused by buyer", "recent",
                                  "missing")}
    k["issued"] = [len(rows), round(sum(r["value"] for r in rows), 0)]
    groups = {}
    for x in dup:
        if x["kind"] == "same amount":
            groups.setdefault(x["eic"], []).append(x)
    k["posted twice"] = [len(groups), round(sum(abs(g[0]["amount"]) * (len(g) - 1) for g in groups.values()), 0)]
    k["bc not in fiscal"] = [len(not_fiscal), round(sum(x["amount"] for x in not_fiscal), 0)]
    k["b2b without eic"] = [len(b2b), round(sum(x["amount"] for x in b2b), 0)]
    k["amount differs net"] = round(sum(r["diff"] or 0 for r in rows if r["status"] == "amount differs"), 0)
    k["refused in bc"] = sum(1 for r in rows if r["status"] == "refused by buyer" and r["bc_no"])
    k["unpaid due"] = [sum(1 for r in rows if r["status_tax"] in ("ACCEPTED", "DELIVERED") and r["due"] and r["due"] < today),
                       round(sum(r["value"] for r in rows if r["status_tax"] in ("ACCEPTED", "DELIVERED") and r["due"]
                                 and r["due"] < today), 0)]
    order = {st: i for i, st in enumerate(["missing", "refused by buyer", "amount differs", "matched, no EIC in BC",
                                           "recent", "matched"])}
    rows.sort(key=lambda r: (order.get(r["status"], 99), -abs(r["value"] or r["total"])))
    return {"kpis": k, "rows": rows, "posted_twice": dup, "bc_not_in_fiscal": not_fiscal, "b2b_without_eic": b2b}


# ---------------------------------------------------------------- the report (background)
def request(client, company, date_from=None, date_to=None, refresh=False):
    s = fiscal.settings()
    cov = coverage(company)
    if not cov["count"] and not cov["loads"]:
        return {"available": False, "coverage": cov, "certificate": cis_client.certificate_info(),
                "error": "No e-invoices issued by us are held yet: read them from the Tax Administration "
                         "(the certificate must be on the server)."}
    b = (date_to or cov["to"] or date.today().isoformat())[:10]
    a = (date_from or max(cov["from"] or b, (date.fromisoformat(b) - timedelta(days=int(fiscal._num(s, "days", 45))))
                          .isoformat()))[:10]
    key = (_key(company), a, b, cov["count"], cov["loads"][0]["at"] if cov["loads"] else "")
    with fiscal._jobs_lock:
        cached = fiscal._results.get(key)
        job = fiscal._jobs.get(_key(company))
        running = job and job["state"] == "running"
    if cached and not refresh and time.time() - cached[0] < 1800:
        return cached[1]
    if not running:
        with fiscal._jobs_lock:
            job = fiscal._jobs[_key(company)] = {"state": "running", "stage": "Business Central", "started": time.time()}
        threading.Thread(target=_run, args=(client, company, a, b, key, job), daemon=True, name="fiscal-sales").start()
    return {"pending": True, "progress": {"stage": job.get("stage")}, "from": a, "to": b,
            "note": "Reading the posted sales invoices and credit memos from Business Central."}


def _run(client, company, a, b, key, job):
    try:
        docs, lines = _bc(client, company, a)
        with fiscal._jobs_lock:
            job["stage"] = "Matching"
        with fiscal._lock:
            cur = fiscal._db().execute("SELECT eic, inv_no, inv_date, buyer_nipt, buyer_name, total, currency, status, "
                                       "due FROM einv_sales WHERE company=? AND inv_date>=? AND inv_date<=?",
                                       (company, a, b))
            cols = [c[0] for c in cur.description]
            einv = [dict(zip(cols, r)) for r in cur.fetchall()]
        out = reconcile(einv, docs, a, b, date.today().isoformat(), fiscal.settings())
        out.update(company=company, date_from=a, date_to=b, checked_at=datetime.now().isoformat(timespec="seconds"),
                   coverage=coverage(company), definitions=DEFINITIONS,
                   bc_read={"lines": lines, "documents": len(docs)})
        with fiscal._jobs_lock:
            fiscal._results[key] = (time.time(), out)
            job["state"] = "done"
    except Exception as exc:                                     # noqa: BLE001
        traceback.print_exc()
        with fiscal._jobs_lock:
            fiscal._results[key] = (time.time(), {"available": False, "error": f"Business Central could not be read: "
                                                                                 f"{str(exc)[:300]}"})
            job["state"] = "failed"


def result_now(client, company, timeout=600):
    t0 = time.time()
    r = request(client, company)
    while r.get("pending") and time.time() - t0 < timeout:
        time.sleep(3)
        r = request(client, company)
    return r


DEFINITIONS = {
    "scope": "E-invoices we issued (to businesses and public bodies), as the Tax Administration lists them for our "
             "NIPT as seller. Sales to individuals at the till are fiscalized by the till and are not in this list.",
    "matched": "Matched: the e-invoice code (EIC) is the NIVF of a posted sales invoice or credit memo and the amounts "
               "agree. Matched, no EIC in BC: same buyer NIPT and amount on a posted document without NIVF.",
    "missing": "Missing: issued at the Tax Administration (for example from the Self-Care portal) and not found in "
               "BC after the grace days. Recent: the same within the grace days.",
    "refused": "Refused by buyer: the buyer refused the e-invoice. If BC still holds it, a credit memo is due.",
    "b2b": "B2B without e-invoice: posted invoices of the period to a customer with an Albanian NIPT that carry no "
           "e-invoice code (our own NIPT excluded).",
    "bc_side": "BC not in fiscal: posted documents whose NIVF the Tax Administration does not list for us. Posted "
               "twice: the same code on two BC documents.",
    "unpaid": "Due and not paid: the Tax Administration status is still Accepted or Delivered after the due date "
              "(the buyer has not registered a payment there).",
}


def render_xlsx(r):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    fill, bold = PatternFill("solid", fgColor="E4F3F9"), Font(bold=True)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = fill, bold
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    k = r["kpis"]
    ws = wb.active
    ws.title = "Summary"
    labels = [("issued", "E-invoices issued in the period"), ("matched", "Matched"),
              ("matched, no EIC in BC", "Matched by NIPT and amount, no EIC in BC"),
              ("amount differs", "Amount differs"), ("refused by buyer", "Refused by the buyer"),
              ("recent", "Recent, within the grace days"), ("missing", "Issued but missing in BC"),
              ("posted twice", "BC: same EIC on two documents, same amount (extra value)"),
              ("bc not in fiscal", "BC: EIC not listed by the Tax Administration"),
              ("b2b without eic", "BC: B2B invoices without e-invoice"),
              ("unpaid due", "Due and not paid at the Tax Administration")]
    sheet(ws, ["Check", "Count", "Value (ALL)"], [[l, k[key][0], k[key][1]] for key, l in labels], [52, 10, 16])
    ws.append([])
    ws.append([f"{r['company']} · e-invoices received by the platform {r['date_from']} to {r['date_to']} · "
               f"checked {r['checked_at']}"])
    sheet(wb.create_sheet("E-invoices issued"),
          ["Status", "Date", "Buyer", "NIPT", "Number", "Total", "Currency", "Tax Administration status", "Due",
           "EIC", "BC document", "BC customer", "BC posting date", "BC amount", "Difference", "Note"],
          [[x["status"], x["date"], x["buyer"], x["nipt"], x["inv_no"], x["total"], x["currency"], x["status_tax"],
            x["due"], x["eic"], x["bc_no"], x["bc_customer"], x["bc_date"], x["bc_amount"], x["diff"], x["note"]]
           for x in r["rows"]], [20, 11, 30, 13, 12, 14, 8, 14, 11, 38, 11, 26, 12, 14, 12, 40])
    sheet(wb.create_sheet("B2B without e-invoice"),
          ["BC document", "Kind", "Customer", "Customer no.", "NIPT", "Posting date", "Amount (ALL)", "Currency",
           "Location", "Description"],
          [[x["bc_no"], x["kind"], x["customer"], x["customer_no"], x["nipt"], x["posting_date"], x["amount"],
            x["currency"], x["location"], x["description"]] for x in r["b2b_without_eic"]],
          [11, 11, 30, 11, 13, 12, 14, 8, 12, 30])
    sheet(wb.create_sheet("BC not in fiscal"), ["BC document", "Kind", "Customer", "NIPT", "EIC", "Posting date",
                                                "Amount (ALL)"],
          [[x["bc_no"], x["kind"], x["customer"], x["nipt"], x["eic"], x["posting_date"], x["amount"]]
           for x in r["bc_not_in_fiscal"]], [11, 11, 30, 13, 38, 12, 14])
    sheet(wb.create_sheet("Posted twice"), ["EIC", "BC document", "Customer", "Posting date", "Amount (ALL)", "Kind"],
          [[x["eic"], x["bc_no"], x["customer"], x["posting_date"], x["amount"], x["kind"]] for x in r["posted_twice"]],
          [38, 11, 30, 12, 14, 16])
    ws = wb.create_sheet("Definitions")
    for key, v in DEFINITIONS.items():
        ws.append([key, v])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_name(r):
    return f"RHG_{r['company'].replace(' ', '_')}_fiscal_sales_vs_BC_{r['date_from']}_{r['date_to']}.xlsx"


def email_block(r):
    """(issues, html, text) for the daily email."""
    import html as H
    k = r["kpis"]
    rows = [("Issued but missing in BC", k["missing"]), ("Refused by the buyer", k["refused by buyer"]),
            ("Amount differs", k["amount differs"]), ("B2B invoices in BC without e-invoice", k["b2b without eic"]),
            ("Same EIC on two BC documents", k["posted twice"]), ("Matched", k["matched"])]
    issues = k["missing"][0] + k["refused by buyer"][0] + k["amount differs"][0] + k["b2b without eic"][0] + \
        k["posted twice"][0]
    trs = "".join(f'<tr><td style="padding:6px 10px;border-bottom:1px solid #e6e9ea">{H.escape(l)}</td>'
                  f'<td style="padding:6px 10px;border-bottom:1px solid #e6e9ea;text-align:right;'
                  f'{"color:#c0392b;font-weight:700" if v[0] and l != "Matched" else ""}">{v[0]}</td>'
                  f'<td style="padding:6px 10px;border-bottom:1px solid #e6e9ea;text-align:right">{fiscal._m(v[1])}</td></tr>'
                  for l, v in rows)
    html_part = (f'<h3 style="margin:22px 0 6px">Sales: e-invoices issued against BC</h3>'
                 f'<p style="color:#5c7482;margin:0 0 8px">{k["issued"][0]} e-invoices, {fiscal._m(k["issued"][1])} ALL, '
                 f'{r["date_from"]} to {r["date_to"]}</p><table style="border-collapse:collapse;width:100%;font-size:14px">'
                 f'<tr style="background:#e4f3f9"><th style="padding:6px 10px;text-align:left">Check</th>'
                 f'<th style="padding:6px 10px;text-align:right">Count</th><th style="padding:6px 10px;text-align:right">'
                 f'Value (ALL)</th></tr>{trs}</table>')
    text = "\nSales, e-invoices issued against BC:\n" + "\n".join(f"{l}: {v[0]} ({fiscal._m(v[1])} ALL)" for l, v in rows)
    return issues, html_part, text
