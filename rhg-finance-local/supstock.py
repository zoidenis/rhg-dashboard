"""Kartela → Suppliers & stock: what we owe each supplier next to the stock still on hand from their deliveries.

For every supplier of the company:
  * payables: their open ledger entries today: owed, overdue, overdue beyond 60 days, not due yet, the oldest
    due date, invoice by invoice;
  * stock on hand from them: the item ledger entries still open (Remaining Quantity above zero) valued at
    their cost (actual + expected) for the quantity left:
      - a purchase receipt still open is theirs exactly: the receipt's invoice (value entries of the invoice
        → the posted invoice → its supplier);
      - stock that came by transfer or a positive adjustment, or a receipt not invoiced yet or older than the
        window read, is split by the item's purchases of the window, supplier by supplier ("allocated");
      - an item never bought in the window stays with "Supplier not known";
  * how fast it moves: sales and consumption of the item in the last 90 days, company-wide (Inventory →
    Movement & dead stock), as days of cover and the stock that did not move at all.
So for a supplier we pay late: what we owe them, and the unsold stock of their products we hold (by category:
wine and spirits turn slowly). A statement per supplier (Excel and PDF) puts both side by side to send. BC is
only read.
"""
import io
import threading
import time
import traceback
from datetime import date, timedelta

import localdb
import movement
import store

_jobs, _lock = {}, threading.Lock()
KEEP = 1800
_hdr = {}
UNKNOWN = "?"


def _f(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _headers(client, company, a, b):
    """{invoice no: (vendor no, vendor name)} of the window, kept an hour."""
    key = (company, a, b)
    hit = _hdr.get(key)
    if hit and time.time() - hit[0] < 3600:
        return hit[1]
    out = {}
    for h in client.purchase_invoice_headers(company, a, b):
        out[str(h.get("No") or "")] = ((h.get("Buy_from_Vendor_No") or "").strip(), (h.get("Buy_from_Vendor_Name") or "").strip())
    _hdr.clear()
    _hdr[key] = (time.time(), out)
    return out


def _usage(client, company, today, progress):
    """{item: quantity sold and consumed in the last 90 days, company-wide}."""
    a = today - timedelta(days=90)
    m = a.replace(day=1)
    while m <= today:
        progress(f"Business Central: item ledger of {m:%B %Y}")
        try:
            movement.load_month(client, company, m.strftime("%Y-%m"))
        except Exception:                                           # noqa: BLE001
            traceback.print_exc()
        m = (m.replace(day=28) + timedelta(days=4)).replace(day=1)
    with movement._lock:
        rows = movement._db().execute(
            "SELECT item, SUM(qty) FROM mv WHERE company=? AND day>=? AND day<=? AND kind IN ('sale','sale_return','consumption') "
            "GROUP BY item", (company, a.isoformat(), today.isoformat())).fetchall()
    return {str(i): -float(q or 0) for i, q in rows}


def build(client, company, months, progress):
    today = date.today()
    w_from = (today.replace(day=1) - timedelta(days=31 * (months - 1))).replace(day=1)
    a, b = w_from.isoformat(), today.isoformat()
    errors = []

    # ---- payables
    progress("Business Central: open supplier entries")
    sup = {}

    def S(no, name=""):
        no = no or UNKNOWN
        x = sup.get(no)
        if x is None:
            x = sup[no] = {"no": no, "name": name or ("Supplier not known" if no == UNKNOWN else no), "owed": 0.0,
                           "overdue": 0.0, "overdue_60": 0.0, "not_due": 0.0, "advance": 0.0, "oldest_due": "",
                           "invoices": [], "stock": 0.0, "exact": 0.0, "allocated": 0.0, "dead": 0.0, "slow": 0.0,
                           "usage_90": 0.0, "items": {}, "cats": {}}
        elif name and x["name"] == no:
            x["name"] = name
        return x
    for e in client.open_vendor_entries(company):
        rem = -_f(e.get("Remaining_Amt_LCY"))
        if abs(rem) < 1:
            continue
        x = S((e.get("Vendor_No") or "").strip(), (e.get("Vendor_Name") or "").strip())
        due = str(e.get("Due_Date") or "")[:10]
        if rem < 0:
            x["advance"] += -rem
            continue
        x["owed"] += rem
        late = (today - date.fromisoformat(due)).days if due > "0001-01-01" else 0
        if late > 0:
            x["overdue"] += rem
            if late > 60:
                x["overdue_60"] += rem
            if not x["oldest_due"] or due < x["oldest_due"]:
                x["oldest_due"] = due
        else:
            x["not_due"] += rem
        x["invoices"].append([str(e.get("Document_No") or ""), str(e.get("Document_Type") or ""),
                              str(e.get("Posting_Date") or "")[:10], due, max(late, 0), round(rem)])

    # ---- stock on hand, entry by entry
    progress("Business Central: stock on hand (open item ledger entries)")
    open_rows = client.open_item_entries(company)
    purchase_ile = {int(r["Entry_No"]) for r in open_rows if r.get("Entry_Type") == "Purchase" and r.get("Entry_No") is not None}
    progress("Business Central: purchase invoices of the window")
    kind = client.hold_purchase_invoice_values(company, a, b)
    heads = _headers(client, company, a, b)
    ile_vendor, item_share = {}, {}
    for v in localdb.iter_rows(kind, company, a, b):
        doc = str(v.get("Document_No") or "")
        vend = heads.get(doc)
        if not vend or not vend[0]:
            continue
        no = int(v.get("Item_Ledger_Entry_No") or 0)
        if no in purchase_ile:
            ile_vendor[no] = vend
        sh = item_share.setdefault(str(v.get("Item_No") or ""), {})
        sh[vend] = sh.get(vend, 0.0) + abs(_f(v.get("Cost_Amount_Actual")))
    progress("Sales and consumption of the last 90 days")
    try:
        usage = _usage(client, company, today, progress)
    except Exception as exc:                                        # noqa: BLE001
        traceback.print_exc()
        errors.append(f"Sales and consumption could not be read: {str(exc)[:160]}")
        usage = {}

    stock_by_item = {}
    for r in open_rows:
        rem, qty = _f(r.get("Remaining_Quantity")), _f(r.get("Quantity"))
        if rem <= 0:
            continue
        cost = _f(r.get("Cost_Amount_Actual")) + _f(r.get("Cost_Amount_Expected"))
        val = cost * rem / qty if qty else 0.0
        item = str(r.get("Item_No") or "")
        it = stock_by_item.setdefault(item, {"qty": 0.0, "value": 0.0})
        it["qty"] += rem
        it["value"] += val
        r["_val"], r["_rem"] = val, rem
    for r in open_rows:
        if "_val" not in r:
            continue
        item, val, rem = str(r.get("Item_No") or ""), r["_val"], r["_rem"]
        exact = ile_vendor.get(int(r.get("Entry_No") or 0)) if r.get("Entry_Type") == "Purchase" else None
        if exact:
            parts = [(exact, 1.0, True)]
        else:
            sh = item_share.get(item) or {}
            tot = sum(sh.values())
            parts = [(vn, v / tot, False) for vn, v in sh.items()] if tot else [((UNKNOWN, ""), 1.0, False)]
        it = stock_by_item[item]
        use = usage.get(item, 0.0)
        cover = (it["qty"] / (use / 90)) if use > 0 else None
        cat = (r.get("Item_Category_Code") or "").strip() or "No category"
        for (vno, vname), share, is_exact in parts:
            x = S(vno, vname)
            v, q = val * share, rem * share
            x["stock"] += v
            x["exact" if is_exact else "allocated"] += v
            x["cats"][cat] = x["cats"].get(cat, 0.0) + v
            if use <= 0:
                x["dead"] += v
            elif cover is not None and cover > 90:
                x["slow"] += v
            I = x["items"].setdefault(item, {"desc": (r.get("Item_Description") or "").strip(), "cat": cat,
                                              "uom": r.get("Unit_of_Measure_Code") or "", "qty": 0.0, "value": 0.0,
                                              "exact": 0.0, "locs": {}, "first": "", "share": 0.0})
            I["qty"] += q
            I["value"] += v
            if is_exact:
                I["exact"] += v
            loc = r.get("Location_Code") or ""
            I["locs"][loc] = I["locs"].get(loc, 0.0) + q
            pd = str(r.get("Posting_Date") or "")[:10]
            if pd and (not I["first"] or pd < I["first"]):
                I["first"] = pd

    # ---- per supplier: the items with their pace, rounding
    out = []
    for x in sup.values():
        items = []
        for item, I in x["items"].items():
            use = usage.get(item, 0.0)
            it = stock_by_item.get(item) or {"qty": 0}
            share = I["qty"] / it["qty"] if it["qty"] else 0
            u90 = use * share
            x["usage_90"] += u90 * (I["value"] / I["qty"] if I["qty"] else 0)
            items.append([item, I["desc"], I["cat"], round(I["qty"], 2), I["uom"], round(I["value"]), round(I["exact"]),
                          sorted(((k, round(v, 2)) for k, v in I["locs"].items() if round(v, 2)), key=lambda kv: -kv[1]),
                          I["first"], (today - date.fromisoformat(I["first"])).days if I["first"] else None,
                          round(u90, 2), round(I["qty"] / (u90 / 90)) if u90 > 0 else None])
        items.sort(key=lambda r: -r[5])
        x["items"] = items
        x["item_count"] = len(items)
        x["cats"] = sorted(([k, round(v)] for k, v in x["cats"].items() if round(v)), key=lambda kv: -kv[1])
        x["invoices"].sort(key=lambda r: (r[3] or "9999", r[0]))
        for k in ("owed", "overdue", "overdue_60", "not_due", "advance", "stock", "exact", "allocated", "dead", "slow"):
            x[k] = round(x[k])
        x["cover"] = round(x["stock"] / (x["usage_90"] / 90)) if x["usage_90"] > 0 else None
        x["usage_90"] = round(x["usage_90"])
        x["oldest_days"] = (today - date.fromisoformat(x["oldest_due"])).days if x["oldest_due"] else 0
        if x["owed"] or x["stock"] or x["advance"]:
            out.append(x)
    out.sort(key=lambda x: -(min(x["owed"], x["stock"]) or x["owed"] or x["stock"]))
    cats = {}
    for x in out:
        for c, v in x["cats"]:
            cats[c] = cats.get(c, 0) + v
    return {"company": company, "as_of": today.isoformat(), "window": {"from": a, "to": b, "months": months},
            "suppliers": out, "categories": sorted(cats.items(), key=lambda kv: -kv[1]), "errors": errors,
            "read_at": time.strftime("%Y-%m-%d %H:%M"), "definitions": DEFINITIONS}


def request(client, company, months=18, refresh=False):
    key = (company or "", months)
    now = time.time()
    with _lock:
        job = _jobs.get(key)
        if job and job["state"] == "running":
            return {"pending": True, "stage": job.get("stage", ""), "elapsed": int(now - job["started"])}
        if job and job["state"] == "done" and not refresh and now - job["finished"] < KEEP:
            return job["result"]
        if job and job["state"] == "failed" and not refresh and now - job["finished"] < 60:
            return {"available": False, "error": job["error"]}
        job = _jobs[key] = {"state": "running", "started": now, "stage": "Starting"}

    def progress(text):
        with _lock:
            job["stage"] = text

    def run():
        try:
            r = build(client, company, months, progress)
            with _lock:
                job.update(state="done", finished=time.time(), result=r)
        except Exception as exc:                                   # noqa: BLE001
            traceback.print_exc()
            with _lock:
                job.update(state="failed", finished=time.time(), error=str(exc)[:300])
    threading.Thread(target=run, daemon=True, name="supstock").start()
    return {"pending": True, "stage": "Starting", "elapsed": 0}


def held(company, months=18):
    with _lock:
        job = _jobs.get((company or "", months))
        return job.get("result") if job and job.get("state") == "done" else None


# ---------------------------------------------------------------- the statement for one supplier
def _supplier(r, no):
    return next((x for x in r["suppliers"] if x["no"] == no), None)


def statement_xlsx(r, no, currency):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    x = _supplier(r, no)
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E4F3F9")
    ws.append([f"{x['name']} ({x['no']}) · payables and stock on hand · {r['company']} · {r['as_of']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Owed", x["owed"], "Overdue", x["overdue"], "Beyond 60 days", x["overdue_60"], "Stock on hand at cost", x["stock"]])
    ws.append([])
    ws.append(["Open invoices"])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["Document", "Type", "Posted", "Due", "Days overdue", f"Open ({currency})"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head
    for inv in x["invoices"]:
        ws.append(inv)
    ws.append(["Total", "", "", "", "", x["owed"]])
    ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append(["Your products in our stock"])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["Item", "Description", "Category", "Quantity", "Unit", f"Value at cost ({currency})", "Locations",
               "Received since", "Days on hand", "Sold and used, 90 days", "Days of cover"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head
    for it in x["items"]:
        ws.append([it[0], it[1], it[2], it[3], it[4], it[5], ", ".join(f"{k or '-'} {v:g}" for k, v in it[7]), it[8], it[9],
                   it[10], it[11] if it[11] is not None else "no sales"])
    ws.append(["Total", "", "", "", "", x["stock"]])
    ws.cell(ws.max_row, 1).font = bold
    for col, w in zip("ABCDEFGHIJK", (12, 38, 16, 11, 8, 16, 30, 14, 12, 18, 13)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0" if abs(c.value) >= 100 or float(c.value).is_integer() else "#,##0.00"
                c.alignment = Alignment(horizontal="right")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def statement_pdf(r, no, currency):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    x = _supplier(r, no)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=f"{x['name']} statement")
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor("#08303f"))
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, leading=15, spaceBefore=8, textColor=colors.HexColor("#0e2733"))
    p = ParagraphStyle("p", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#5c7482"))
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)

    def n(v, d=0):
        return "–" if v is None else f"{v:,.{d}f}"

    def tbl(rows, widths, right):
        t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
        st = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
              ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e4f3f9")),
              ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#d9dcdd")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for c in right:
            st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
        t.setStyle(TableStyle(st))
        return t
    story = [Paragraph(f"{x['name']} · statement of account and stock on hand", h1),
             Paragraph(f"{r['company']} · {x['no']} · as of {r['as_of']} · amounts in {currency}", p), Spacer(1, 6),
             tbl([["Owed", "Overdue", "Beyond 60 days", "Not due yet", "Your products in our stock (at cost)"],
                  [n(x["owed"]), n(x["overdue"]), n(x["overdue_60"]), n(x["not_due"]), n(x["stock"])]],
                 (40, 40, 40, 40, 80), (0, 1, 2, 3, 4)),
             Paragraph("Open invoices", h2)]
    rows = [["Document", "Type", "Posted", "Due", "Days overdue", "Open"]] + \
        [[i[0], i[1], i[2], i[3], n(i[4]) if i[4] else "", n(i[5])] for i in x["invoices"]] + [["Total", "", "", "", "", n(x["owed"])]]
    story.append(tbl(rows, (34, 30, 26, 26, 26, 34), (4, 5)))
    story.append(Paragraph("Your products in our stock", h2))
    rows = [["Item", "Description", "Category", "Quantity", "Unit", "Value at cost", "Since", "Days held", "Used, 90 d", "Cover, days"]] + \
        [[i[0], Paragraph(i[1], cell), i[2], n(i[3], 2 if i[3] % 1 else 0), i[4], n(i[5]), i[8], n(i[9]), n(i[10], 2 if (i[10] or 0) % 1 else 0),
          n(i[11]) if i[11] is not None else "no sales"] for i in x["items"]] + [["Total", "", "", "", "", n(x["stock"]), "", "", "", ""]]
    story.append(tbl(rows, (18, 70, 26, 20, 12, 26, 24, 20, 26, 22), (3, 5, 7, 8, 9)))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Stock on hand = what is left of each delivery in our stores, at its purchase cost, including what "
                           "moved from our warehouse to the restaurants (attributed by your share of our purchases of the item). "
                           "Since = the oldest delivery still on hand. Cover = the quantity on hand against what was sold and used "
                           "in the last 90 days.", p))
    doc.build(story)
    return buf.getvalue()


DEFINITIONS = {
    "payables": "Payables: the supplier's open ledger entries today. Owed = invoices and other debts still open; "
                "overdue = past the due date; advances = payments not applied yet (shown apart).",
    "stock": "Stock on hand: item ledger entries still open (quantity left above zero) at their cost (actual and "
             "expected) for the quantity left. A purchase receipt still open belongs to the supplier of its invoice "
             "(exact). Stock that came by transfer from the warehouse, a positive adjustment, a receipt not invoiced "
             "yet or older than the window is split among the item's suppliers by their purchases of the window "
             "(allocated). Items not bought in the window: Supplier not known.",
    "pace": "Pace: sales and consumption of the item in the last 90 days, company-wide (Inventory → Movement & dead "
            "stock), shared by the supplier's part of the stock. Days of cover = stock ÷ daily use; not moved = no "
            "sale or consumption in 90 days; slow = more than 90 days of cover.",
}
