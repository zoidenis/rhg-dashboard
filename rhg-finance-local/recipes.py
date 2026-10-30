"""Recipe integrity and item master control (Inventory Control, phase 2).

What Business Central holds (verified on SALT, see the project's Inventory prompt v2):
- Item cards (PbItems): Assembly_Policy, Assembly_BOM, Replenishment_System, categories,
  posting groups, units, costs, Blocked.
- Recipes are Assembly BOMs (AssemblyBOM): one BOM per item, not per location, with no
  modification date.
- A sale of an Assemble-to-Order item posts Sale + Assembly Output + Assembly Consumption,
  the consumption generated from the BOM. An item sold with no BOM is therefore sold with
  no consumption at all, and a broken BOM consumes the wrong things.

Which items need a recipe (rules, not BC fields; shown on screen):
- Assemble-to-Order, or Replenishment System = Assembly, or an item that has a BOM, or a
  PN item (item category set in Settings, default PRODUKT I NDERMJETEM: an assumption until
  Finance confirms it).
- Everything else that sells is treated as sold as purchased (Not required), and says so.

Coverage per location is about the items sold there in the period, because the BOM itself
does not vary by location. Sales come from the item ledger (Entry_Type = Sale), read once
per month into data/recipes.db and refreshed while the month can still change.

Nothing is written to Business Central.
"""
import io
import json
import sqlite3
import threading
import time
import traceback
from calendar import monthrange
from datetime import date, datetime, timedelta

import config
import store

DB_PATH = store.DATA_DIR / "recipes.db"
MASTER_FILE = store.DATA_DIR / "recipes_master.json"
MASTER_TTL = 6 * 3600
OPEN_DAYS = 35
ITEM_FIELDS = ("No,Description,Type,Assembly_Policy,Assembly_BOM,Replenishment_System,Base_Unit_of_Measure,"
               "Sales_Unit_of_Measure,Purch_Unit_of_Measure,Costing_Method,Unit_Cost,Standard_Cost,"
               "Last_Direct_Cost,Item_Category_Code,Item_Category_Group_1,Item_Category_Group_2,"
               "Item_Category_Group_3,Item_Category_Group_4,Inventory_Posting_Group,Gen_Prod_Posting_Group,Blocked,"
               "Last_Date_Modified")
STATUSES = ("VALID", "MISSING", "INCOMPLETE", "INACTIVE", "PENDING REVIEW", "NOT REQUIRED")

_lock = threading.Lock()
_conn = None
_jobs, _jobs_lock = {}, threading.Lock()
_results = {}


# ---------------------------------------------------------------- storage
def _db():
    global _conn
    if _conn is None:
        DB_PATH.parent.mkdir(exist_ok=True)
        _conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        _conn.execute("PRAGMA journal_mode=WAL")
        _conn.executescript("""
            CREATE TABLE IF NOT EXISTS sales (company TEXT, day TEXT, location TEXT, item TEXT,
                                              qty REAL, entries INTEGER);
            CREATE INDEX IF NOT EXISTS sales_day ON sales (company, day);
            CREATE TABLE IF NOT EXISTS loaded (company TEXT, month TEXT, date_to TEXT, at REAL,
                                               PRIMARY KEY (company, month));
        """)
    return _conn


def _master_cache():
    try:
        return json.loads(MASTER_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def master_cached(company):
    """The item cards last read, whatever their age (no read from BC)."""
    with _lock:
        return _master_cache().get(company)


def master(client, company, force=False):
    """Item cards, BOM lines and units of measure, refreshed every six hours."""
    with _lock:
        cache = _master_cache()
    hit = cache.get(company)
    if hit and not force and time.time() - hit.get("at", 0) < MASTER_TTL and hit.get("fields") == ITEM_FIELDS:
        return hit
    items = client._get_all(client._url(company, config.E_ITEMS), {"$select": ITEM_FIELDS})
    bom = client._get_all(client._url(company, "AssemblyBOM"),
                          {"$select": "Parent_Item_No,Line_No,Type,No,Description,Assembly_BOM,Quantity_per,"
                                      "Unit_of_Measure_Code"})
    uom = client._get_all(client._url(company, "ItemUnitsOfMeasure"),
                          {"$select": "Item_No,Code,Qty_per_Unit_of_Measure"})
    value = {"at": time.time(), "read_at": datetime.now().isoformat(timespec="seconds"), "fields": ITEM_FIELDS,
             "items": {str(r["No"]): {k: r.get(k) for k in ITEM_FIELDS.split(",") if k != "No"} for r in items
                       if r.get("No")},
             "bom": [[str(r.get("Parent_Item_No")), r.get("Line_No"), r.get("Type") or "", str(r.get("No") or ""),
                      r.get("Description") or "", bool(r.get("Assembly_BOM")), float(r.get("Quantity_per") or 0),
                      r.get("Unit_of_Measure_Code") or ""] for r in bom],
             "uom": [[str(r.get("Item_No")), str(r.get("Code") or "").upper(),
                      float(r.get("Qty_per_Unit_of_Measure") or 0)] for r in uom]}
    with _lock:
        cache = _master_cache()
        cache[company] = value
        tmp = MASTER_FILE.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(cache, separators=(",", ":")), encoding="utf-8")
        tmp.replace(MASTER_FILE)
    return value


def _month_bounds(month, today):
    y, m = int(month[:4]), int(month[5:7])
    return date(y, m, 1), min(date(y, m, monthrange(y, m)[1]), today)


def _fresh(company, month, b, today, force):
    with _lock:
        row = _db().execute("SELECT date_to, at FROM loaded WHERE company=? AND month=?", (company, month)).fetchone()
    if not row:
        return False
    date_to, at = row
    if date_to < b.isoformat():
        return False
    age = time.time() - at
    open_month = (today - b).days <= OPEN_DAYS
    if force and open_month and age > 120:
        return False
    return age < (3600 if b >= today - timedelta(days=1) else (6 * 3600 if open_month else 30 * 86400))


def load_sales_month(client, company, month, force=False, today=None):
    today = today or date.today()
    a, b = _month_bounds(month, today)
    if a > today or _fresh(company, month, b, today, force):
        return False
    agg = {}

    def page(rows):
        for r in rows:
            k = (str(r.get("Posting_Date"))[:10], r.get("Location_Code") or "", str(r.get("Item_No") or ""))
            x = agg.setdefault(k, [0.0, 0])
            x[0] += -float(r.get("Quantity") or 0)          # sales are negative quantities
            x[1] += 1
    client._get_all(client._url(company, config.E_ILE),
                    {"$filter": f"Entry_Type eq 'Sale' and Posting_Date ge {a.isoformat()} "
                                f"and Posting_Date le {b.isoformat()}",
                     "$select": "Entry_No,Item_No,Location_Code,Posting_Date,Quantity"}, on_page=page)
    with _lock:
        db = _db()
        last = date(a.year, a.month, monthrange(a.year, a.month)[1]).isoformat()
        db.execute("DELETE FROM sales WHERE company=? AND day>=? AND day<=?", (company, a.isoformat(), last))
        db.executemany("INSERT INTO sales VALUES (?,?,?,?,?,?)",
                       [(company, *k, round(v[0], 4), v[1]) for k, v in agg.items()])
        db.execute("INSERT OR REPLACE INTO loaded VALUES (?,?,?,?)", (company, month, b.isoformat(), time.time()))
        db.commit()
    return True


def months_between(a, b):
    out, y, m = [], a.year, a.month
    while (y, m) <= (b.year, b.month):
        out.append(f"{y}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def sales(company, a, b):
    """{(location, item): [qty, entries, last day]}"""
    with _lock:
        rows = _db().execute("SELECT location, item, SUM(qty), SUM(entries), MAX(day) FROM sales "
                             "WHERE company=? AND day>=? AND day<=? GROUP BY location, item",
                             (company, a.isoformat(), b.isoformat())).fetchall()
    return {(r[0], r[1]): [r[2] or 0, r[3] or 0, r[4]] for r in rows}


# ---------------------------------------------------------------- the analysis
def technical_matcher(settings):
    """POS technical items (course markers, price fillers, modifiers) that are not products.
    Settings → inventory.pos_technical_items: item numbers or descriptions; a trailing * means
    'starts with'. Case does not matter."""
    text = (settings.get("inventory") or {}).get("pos_technical_items") or ""
    exact, prefix = set(), []
    for part in text.split(","):
        part = part.strip().upper()
        if not part:
            continue
        (prefix.append(part[:-1].strip()) if part.endswith("*") else exact.add(part))
    prefix = tuple(p for p in prefix if p)

    def match(item, desc):
        i, d = (item or "").upper(), (desc or "").strip().upper()
        return i in exact or d in exact or (bool(prefix) and (i.startswith(prefix) or d.startswith(prefix)))
    return match


def _pn_categories(settings):
    text = (settings.get("inventory") or {}).get("pn_categories") or "PRODUKT I NDERMJETEM"
    return {c.strip().upper() for c in text.split(",") if c.strip()}


def analyse(m, sold, settings, period):
    items, uom = m["items"], {}
    for it, code, per in m["uom"]:
        uom[(it, code)] = per
    lines = {}
    for parent, line_no, typ, no, desc, nested, qty, unit in m["bom"]:
        lines.setdefault(parent, []).append({"line": line_no, "type": typ, "no": no, "desc": desc,
                                             "nested": nested, "qty": qty, "unit": unit.upper()})
    pn = _pn_categories(settings)
    technical = technical_matcher(settings)

    def per_base(item, unit):
        card = items.get(item) or {}
        base = str(card.get("Base_Unit_of_Measure") or "").upper()
        if not unit or unit == base:
            return 1.0
        return uom.get((item, unit))

    def needs_own_recipe(card):
        """A component flagged as an assembly whose BOM has no item lines only breaks the parent recipe
        when it cannot come from stock at a real cost: Assemble-to-Order, a PN item, or no unit cost.
        A bought ingredient (e.g. a bitter set to Replenishment Assembly by mistake) is consumed from
        stock at its own cost, so the parent stays valid; Item master control lists the setup instead."""
        return (card.get("Assembly_Policy") == "Assemble-to-Order"
                or str(card.get("Item_Category_Code") or "").upper() in pn
                or float(card.get("Unit_Cost") or 0) <= 0)

    def has_item_lines(item):
        return any(x["type"] == "Item" for x in lines.get(item) or [])

    memo = {}

    def check(item, stack=()):
        """(status, issues, components, cost) of one item's recipe."""
        if item in memo:
            return memo[item]
        if item in stack:
            return ("INCOMPLETE", [f"circular reference through {item}"], 0, None)
        card = items.get(item) or {}
        bl = lines.get(item) or []
        comps = [l for l in bl if l["type"] == "Item"]
        if not comps:
            res = ("MISSING", ["no Assembly BOM"] if not bl else ["the Assembly BOM has no item lines"], 0, None)
            memo[item] = res
            return res
        issues, cost, blocked = [], 0.0, 0
        for l in comps:
            comp = items.get(l["no"])
            if comp is None:
                issues.append("a BOM line has no item number" if not l["no"]
                              else f"component {l['no']} is not in the item list")
                cost = None
                continue
            if l["qty"] <= 0:
                issues.append(f"{l['no']} {comp.get('Description') or ''}: quantity is zero")
            per = per_base(l["no"], l["unit"])
            if per is None or per <= 0:
                issues.append(f"{l['no']}: unit {l['unit']} has no conversion to {comp.get('Base_Unit_of_Measure')}")
                per = None
            if comp.get("Blocked"):
                blocked += 1
                issues.append(f"{l['no']} {comp.get('Description') or ''}: component is blocked")
            if [x for x in lines.get(l["no"]) or [] if x["type"] == "Item"]:
                sub = check(l["no"], stack + (item,))
                if sub[0] in ("MISSING", "INCOMPLETE"):
                    issues.append(f"{l['no']}: sub-recipe {sub[0].lower()}")
            elif (l["nested"] or lines.get(l["no"])) and needs_own_recipe(comp):
                issues.append(f"{l['no']}: sub-recipe missing (made in-house or no cost, and its BOM has no item lines)")
            if cost is not None and per is not None:
                cost += l["qty"] * per * float(comp.get("Unit_Cost") or 0)
        if card.get("Blocked") or blocked == len(comps):
            status = "INACTIVE"
            issues.insert(0, "the item is blocked" if card.get("Blocked") else "every component is blocked")
        elif issues:
            status = "INCOMPLETE"
        else:
            status = "VALID"
        res = (status, issues, len(comps), round(cost, 2) if cost is not None else None)
        memo[item] = res
        return res

    def requirement(item):
        card = items.get(item)
        if card is None:
            return "PENDING REVIEW", "the item is not in the item list"
        policy = card.get("Assembly_Policy") or ""
        if technical(item, card.get("Description")):
            return "NOT REQUIRED", "POS technical item (list in Settings), not a product"
        if policy == "Assemble-to-Order":
            return "required", "Assemble-to-Order"
        if (card.get("Replenishment_System") or "") == "Assembly":
            return "required", "replenished by assembly"
        if lines.get(item):
            return "required", "has an Assembly BOM"
        if str(card.get("Item_Category_Code") or "").upper() in pn:
            return "required", "PN item (category rule in Settings)"
        if float(card.get("Unit_Cost") or 0) <= 0:
            return "PENDING REVIEW", "set up as sold as purchased but has no unit cost: with no recipe and " \
                                     "no cost it sells with no consumption (a combo or menu item without a recipe?)"
        return "NOT REQUIRED", f"sold as purchased ({policy or 'no policy'}, replenished by " \
                               f"{card.get('Replenishment_System') or 'purchase'})"

    def status_of(item):
        req, why = requirement(item)
        if req in ("PENDING REVIEW", "NOT REQUIRED"):
            return req, why, [], 0, None
        st, issues, n, cost = check(item)
        return st, why, issues, n, cost

    # ---- items sold in the period, per location and in total
    sold_items = {}
    for (loc, it), (q, n, last) in sold.items():
        t = sold_items.setdefault(it, [0.0, 0, "", set()])
        t[0] += q
        t[1] += n
        t[2] = max(t[2], last or "")
        t[3].add(loc)
    info = {}
    for it in set(sold_items) | {k for k, v in items.items() if v.get("Assembly_Policy") == "Assemble-to-Order"}:
        st, why, issues, n, cost = status_of(it)
        card = items.get(it) or {}
        s = sold_items.get(it)
        info[it] = {"item": it, "desc": card.get("Description") or "", "status": st, "why": why, "issues": issues,
                    "components": n, "recipe_cost": cost, "policy": card.get("Assembly_Policy") or "",
                    "category": card.get("Item_Category_Code") or "", "group1": card.get("Item_Category_Group_1") or "",
                    "pn": str(card.get("Item_Category_Code") or "").upper() in pn,
                    "blocked": bool(card.get("Blocked")), "unit_cost": card.get("Unit_Cost"),
                    "qty_sold": round(s[0], 2) if s else 0, "sale_entries": s[1] if s else 0,
                    "last_sale": s[2] if s else "", "locations": sorted(s[3]) if s else []}

    # ---- coverage per location: the items sold there
    by_loc = {}
    for (loc, it), (q, n, last) in sold.items():
        st = info[it]["status"]
        L = by_loc.setdefault(loc, {"location": loc, "sold": 0, "required": 0, "rows": []})
        L["sold"] += 1
        L.setdefault(st, 0)
        L[st] += 1
        if st not in ("NOT REQUIRED",):
            L["required"] += 1 if st != "PENDING REVIEW" else 0
        L["rows"].append([it, st, round(q, 2), n, last])
    coverage = []
    for loc, L in sorted(by_loc.items()):
        req = L["required"]
        coverage.append({"location": loc or "(none)", "sold_items": L["sold"], "required": req,
                         **{s: L.get(s, 0) for s in STATUSES},
                         "coverage": round(L.get("VALID", 0) / req * 100, 1) if req else None, "rows": L["rows"]})
    tot = {s: sum(1 for i in info.values() if i["sale_entries"] and i["status"] == s) for s in STATUSES}
    tot_req = sum(tot[s] for s in ("VALID", "MISSING", "INCOMPLETE", "INACTIVE"))

    # ---- the ATO items with no BOM, sold or not
    missing_all = sorted((i for i in info.values() if i["status"] == "MISSING"),
                         key=lambda i: (-i["qty_sold"], i["item"]))
    incomplete = sorted((i for i in info.values() if i["status"] in ("INCOMPLETE", "INACTIVE")),
                        key=lambda i: (-i["qty_sold"], i["item"]))
    pending = sorted((i for i in info.values() if i["status"] == "PENDING REVIEW" and i["sale_entries"]),
                     key=lambda i: (-i["qty_sold"], i["item"]))

    # ---- item master control
    used_in_sold = {}
    for it, i in info.items():
        if i["qty_sold"]:
            for l in lines.get(it, []):
                if l["type"] == "Item":
                    used_in_sold.setdefault(l["no"], set()).add(it)
    checks = {k: [] for k in ("no_category", "no_inventory_group", "no_product_group", "no_base_unit",
                               "unit_no_conversion", "zero_cost_in_use", "cost_far_from_last_direct",
                               "blocked_sold", "blocked_in_recipes", "ato_without_bom",
                               "assembly_setup_but_bought")}
    for it, card in items.items():
        active = not card.get("Blocked") and card.get("Type") == "Inventory"
        desc = card.get("Description") or ""
        sold_q = sold_items.get(it, [0])[0]
        in_use = bool(sold_q) or it in used_in_sold
        base = str(card.get("Base_Unit_of_Measure") or "").upper()
        if active:
            if not card.get("Item_Category_Code"):
                checks["no_category"].append([it, desc, ""])
            if not card.get("Inventory_Posting_Group"):
                checks["no_inventory_group"].append([it, desc, ""])
            if not card.get("Gen_Prod_Posting_Group"):
                checks["no_product_group"].append([it, desc, ""])
            if not base:
                checks["no_base_unit"].append([it, desc, ""])
            else:
                for label, f in (("purchase", "Purch_Unit_of_Measure"), ("sales", "Sales_Unit_of_Measure")):
                    u = str(card.get(f) or "").upper()
                    if u and u != base and (it, u) not in uom:
                        checks["unit_no_conversion"].append([it, desc, f"{label} unit {u}, base {base}"])
            uc, ld = float(card.get("Unit_Cost") or 0), float(card.get("Last_Direct_Cost") or 0)
            is_tech = technical(it, desc)
            if in_use and uc <= 0 and card.get("Assembly_Policy") != "Assemble-to-Order" and not is_tech:
                checks["zero_cost_in_use"].append([it, desc, "sold" if sold_q else "used in a sold recipe"])
            if (card.get("Replenishment_System") or "") == "Purchase" and uc > 0 and ld > 0 and \
                    (uc > ld * 3 or uc < ld / 3):
                checks["cost_far_from_last_direct"].append([it, desc, f"unit cost {uc:,.2f}, last direct {ld:,.2f}"])
            if card.get("Assembly_Policy") == "Assemble-to-Order" and not lines.get(it) and not is_tech:
                checks["ato_without_bom"].append([it, desc, f"sold {sold_q:,.0f}" if sold_q else "not sold"])
            if card.get("Assembly_Policy") != "Assemble-to-Order" and not is_tech and not has_item_lines(it) and \
                    ((card.get("Replenishment_System") or "") == "Assembly" or card.get("Assembly_BOM")) and ld > 0:
                where = f"in {len(used_in_sold[it])} sold recipe(s)" if it in used_in_sold else \
                    ("sold" if sold_q else "not in use")
                checks["assembly_setup_but_bought"].append(
                    [it, desc, f"replenished by {card.get('Replenishment_System') or '-'}, "
                               f"{len(lines.get(it) or [])} BOM line(s) without an item, last direct cost "
                               f"{ld:,.2f}; {where}"])
        if card.get("Blocked"):
            if sold_q:
                checks["blocked_sold"].append([it, desc, f"sold {sold_q:,.0f} in the period"])
            if it in used_in_sold:
                parents = sorted(used_in_sold[it])
                checks["blocked_in_recipes"].append([it, desc, f"in {len(parents)} sold recipe(s): "
                                                     + ", ".join(parents[:5]) + (" …" if len(parents) > 5 else "")])
    counts = {k: len(v) for k, v in checks.items()}
    active = [c for c in items.values() if not c.get("Blocked") and c.get("Type") == "Inventory"]
    names = {it: (items.get(it) or {}).get("Description") or "" for L in coverage for it, *_ in L["rows"]}
    return {
        "period": period,
        "kpis": {"items": len(items), "active_items": len(active),
                 "ato": sum(1 for c in items.values() if c.get("Assembly_Policy") == "Assemble-to-Order"),
                 "ato_without_bom": counts["ato_without_bom"],
                 "pn": sum(1 for c in items.values() if str(c.get("Item_Category_Code") or "").upper() in pn),
                 "with_bom": len(lines), "sold_items": len(sold_items),
                 "required": tot_req, **{s.lower().replace(" ", "_"): tot[s] for s in STATUSES},
                 "coverage": round(tot["VALID"] / tot_req * 100, 1) if tot_req else None,
                 "missing_sold_qty": round(sum(i["qty_sold"] for i in missing_all), 0),
                 "missing_sale_entries": sum(i["sale_entries"] for i in missing_all),
                 "incomplete_sold_qty": round(sum(i["qty_sold"] for i in incomplete), 0)},
        "coverage": coverage, "missing": missing_all, "incomplete": incomplete, "review": pending,
        "checks": {k: v[:1000] for k, v in checks.items()}, "check_counts": counts,
        "names": names, "pn_rule": sorted(pn),
        "master_read_at": m.get("read_at"),
    }


# ---------------------------------------------------------------- the request
def parse_range(date_from, date_to, today=None):
    today = today or date.today()
    if not date_from or not date_to:
        b = today - timedelta(days=1)
        return b - timedelta(days=29), b
    a, b = date.fromisoformat(date_from[:10]), min(date.fromisoformat(date_to[:10]), today)
    if a > b:
        raise ValueError("The start date is after the end date.")
    if (b - a).days > 366:
        raise ValueError("Choose a range of one year or less.")
    return a, b


def request(client, company, date_from, date_to, refresh=False):
    a, b = parse_range(date_from, date_to)
    key = (company, a.isoformat(), b.isoformat())
    now = time.time()
    with _jobs_lock:
        job = _jobs.get(company)
        running = job and job["state"] == "running"
        cached = _results.get(key)
    if cached and not refresh and now - cached[0] < 900 and not running:
        return cached[1]
    if not running:
        with _jobs_lock:
            job = _jobs[company] = {"state": "running", "stage": "Item cards and recipes", "done": 0,
                                    "total": 0, "key": key, "started": now, "errors": []}
        threading.Thread(target=_run, args=(client, company, a, b, refresh, job), daemon=True,
                         name="recipes").start()
    with _jobs_lock:
        snap = dict(_jobs[company])
    return {"pending": True, "progress": {k: snap.get(k) for k in ("stage", "done", "total")},
            "period": {"from": a.isoformat(), "to": b.isoformat()},
            "note": "Reading the item cards, the Assembly BOMs and the sales of the period from Business "
                    "Central. The first read takes a few minutes; after that it is quick."}


def _run(client, company, a, b, refresh, job):
    try:
        m = master(client, company, force=refresh)
        months = months_between(a, b)
        with _jobs_lock:
            job.update(stage="Sales by item and location", total=len(months))
        for mo in reversed(months):
            try:
                load_sales_month(client, company, mo, force=refresh)
            except Exception as exc:                               # noqa: BLE001
                job["errors"].append(f"sales {mo}: {str(exc)[:200]}")
            with _jobs_lock:
                job["done"] += 1
        with _jobs_lock:
            job.update(stage="Checking recipes")
        out = analyse(m, sales(company, a, b), store.get_settings(), {"from": a.isoformat(), "to": b.isoformat()})
        out.update(company=company, errors=job["errors"], checked_at=datetime.now().isoformat(timespec="seconds"),
                   definitions=DEFINITIONS)
        with _jobs_lock:
            _results[(company, a.isoformat(), b.isoformat())] = (time.time(), out)
            if len(_results) > 6:
                for k, _ in sorted(_results.items(), key=lambda kv: kv[1][0])[:2]:
                    _results.pop(k, None)
            job["state"] = "done"
    except Exception as exc:                                       # noqa: BLE001
        traceback.print_exc()
        with _jobs_lock:
            job.update(state="failed", errors=job["errors"] + [str(exc)[:300]])
            _results[(company, a.isoformat(), b.isoformat())] = (time.time(), {
                "available": False, "error": f"The recipes could not be read: {str(exc)[:300]}",
                "period": {"from": a.isoformat(), "to": b.isoformat()}})


# ---------------------------------------------------------------- Excel
def render_xlsx(r):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    fill, bold = PatternFill("solid", fgColor="E4F3F9"), Font(bold=True)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = fill, bold
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    ws = wb.active
    ws.title = "Coverage"
    sheet(ws, ["Location", "Items sold", "Recipe required", "Valid", "Missing", "Incomplete", "Inactive",
               "Pending review", "Not required", "Coverage %"],
          [[c["location"], c["sold_items"], c["required"], c["VALID"], c["MISSING"], c["INCOMPLETE"], c["INACTIVE"],
            c["PENDING REVIEW"], c["NOT REQUIRED"], c["coverage"]] for c in r["coverage"]],
          [16, 11, 15, 9, 9, 11, 9, 14, 12, 11])

    def item_rows(lst):
        return [[i["item"], i["desc"], i["policy"], i["category"], "yes" if i["pn"] else "", i["status"],
                 "; ".join(i["issues"]), i["components"], i["recipe_cost"], i["qty_sold"], i["sale_entries"],
                 i["last_sale"], ", ".join(i["locations"])] for i in lst]
    head = ["Item", "Description", "Assembly policy", "Category", "PN", "Status", "Issues", "Components",
            "Recipe cost (derived)", "Qty sold", "Sale entries", "Last sale", "Locations"]
    widths = [12, 40, 18, 16, 5, 12, 60, 11, 16, 10, 11, 11, 30]
    sheet(wb.create_sheet("Missing recipes"), head, item_rows(r["missing"]), widths)
    sheet(wb.create_sheet("Incomplete & inactive"), head, item_rows(r["incomplete"]), widths)
    sheet(wb.create_sheet("Pending review"), head + ["Reason"],
          [row + [i["why"]] for row, i in zip(item_rows(r.get("review") or []), r.get("review") or [])],
          widths + [60])
    rows = []
    for loc in r["coverage"]:
        for it, st, q, n, last in loc["rows"]:
            rows.append([loc["location"], it, r["names"].get(it, ""), st, q, n, last])
    sheet(wb.create_sheet("Sold items by location"), ["Location", "Item", "Description", "Recipe status", "Qty sold",
                                                      "Sale entries", "Last sale"], rows, [14, 12, 40, 14, 10, 11, 11])
    labels = CHECK_LABELS
    rows = [[labels.get(k, k), it, desc, note] for k, lst in r["checks"].items() for it, desc, note in lst]
    sheet(wb.create_sheet("Item master checks"), ["Check", "Item", "Description", "Detail"], rows, [40, 12, 40, 50])
    ws = wb.create_sheet("Definitions")
    for k, v in DEFINITIONS.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CHECK_LABELS = {
    "ato_without_bom": "Assemble-to-Order without an Assembly BOM",
    "assembly_setup_but_bought": "Set up as assembled but bought (BOM has no item lines)",
    "blocked_sold": "Blocked item still sold",
    "blocked_in_recipes": "Blocked item used in a sold recipe",
    "zero_cost_in_use": "Zero unit cost on an item in use",
    "cost_far_from_last_direct": "Unit cost over 3x away from the last direct cost",
    "unit_no_conversion": "Purchase or sales unit without a conversion",
    "no_base_unit": "No base unit of measure",
    "no_category": "No item category",
    "no_inventory_group": "No inventory posting group",
    "no_product_group": "No general product posting group",
}

DEFINITIONS = {
    "required": "A recipe is required for Assemble-to-Order items, items replenished by assembly, items that have "
                "an Assembly BOM, and PN items (item category rule in Settings). Other items that sell are "
                "treated as sold as purchased (Not required); the reason is shown on each item.",
    "valid": "VALID: every component exists, is not blocked, has a quantity above zero and a unit that converts "
             "to the component's base unit, and every sub-recipe resolves.",
    "missing": "MISSING: a recipe is required and the item has no Assembly BOM (or the BOM has no item lines). "
               "In Business Central such an item sells with no ingredient consumption at all.",
    "incomplete": "INCOMPLETE: the BOM exists but has a zero quantity, a missing or blocked component, a unit "
                  "without conversion, a broken sub-recipe or a circular reference. A component set up as an "
                  "assembly with no item lines in its BOM counts as a broken sub-recipe only when it is "
                  "Assemble-to-Order, a PN item or has no unit cost; a bought ingredient is consumed from stock.",
    "inactive": "INACTIVE: the item itself is blocked, or every component is blocked.",
    "pending": "PENDING REVIEW: the item sold but is not in the item list as read, or it is set up as sold as "
               "purchased (no recipe needed) yet has no unit cost, which usually means a combo or menu item "
               "that is missing its recipe.",
    "technical": "POS technical items listed in Settings (course markers, price fillers, modifiers) are not "
                 "products and need no recipe.",
    "coverage": "Recipe coverage % = valid / (items sold that require a recipe) x 100, per location on the items "
                "sold at that location. Business Central keeps one BOM per item, not per location.",
    "cost": "Recipe cost is derived: sum of quantity per x unit conversion x the component's current unit cost. "
            "It is not a posted value.",
    "sales": "Sales are item ledger entries of type Sale in the period, by location; quantities in base units.",
    "not_available": "Not available in Business Central: the date a recipe was last changed, the Simphony Menu Item "
                     "Number (mapping to be verified), and recipes by location.",
}
