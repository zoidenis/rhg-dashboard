"""Fiscalized purchase invoices against Business Central (the Self-Care reconciliation).

Every purchase invoice a supplier issues to us is fiscalized at the Tax Administration and
shows in Self-Care with its e-invoice code (EIC). Business Central keeps the same code on
the posted purchase invoice (field NIVF). This module puts the two side by side:

Sources of the fiscalized invoices (kept in data/fiscal.db, newest wins per EIC):
- the e-invoice service of the Tax Administration, read every morning with the
  fiscalization certificate (cis_client.py), the automatic way;
- the Self-Care purchases book ("Libri i Blerjeve", .xlsx), imported by hand, for months
  before the service was connected or when it is down.

Matching, in this order, each BC document used once:
1. EIC = NIVF on a posted purchase invoice; amounts compared in lek (a foreign-currency
   invoice is compared at the lek amount BC posted to the supplier ledger).
2. Supplier NIPT + amount (within the tolerance) on a posted invoice without that code,
   invoice dates within 45 days, the closest first: BC lacks the NIVF.
3. Credit notes (negative totals): supplier NIPT + amount on a posted credit memo.
4. The EIC on a purchase order that is not posted yet: entered, not booked.
What is left is missing in BC, or "recent" while younger than the grace days.

From the BC side: the same EIC on two posted invoices (same amount: possibly posted twice),
posted invoices whose EIC the Tax Administration does not list for the period, and posted
invoices of the period without an EIC.

Business Central is only read. Nothing is sent to the Tax Administration but reading
requests (see cis_client.ALLOWED).
"""
import base64
import io
import json
import re
import sqlite3
import threading
import time
import traceback
import uuid
from datetime import date, datetime, timedelta

import arap_notify
import cis_client
import config
import store

DB_PATH = store.DATA_DIR / "fiscal.db"
STATE_FILE = store.DATA_DIR / "fiscal_state.json"
_lock = threading.Lock()
_conn = None
_jobs, _jobs_lock = {}, threading.Lock()
_results = {}
_state_lock = threading.Lock()
_run_lock = threading.Lock()


def _db():
    global _conn
    if _conn is None:
        DB_PATH.parent.mkdir(exist_ok=True)
        _conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        _conn.execute("PRAGMA journal_mode=WAL")
        _conn.executescript("""
            CREATE TABLE IF NOT EXISTS einv (company TEXT, eic TEXT, iic TEXT, inv_no TEXT, inv_date TEXT,
                seller_nipt TEXT, seller_name TEXT, buyer_nipt TEXT, total REAL, currency TEXT, doc_type TEXT,
                status TEXT, source TEXT, seen_at TEXT, PRIMARY KEY (company, eic));
            CREATE INDEX IF NOT EXISTS einv_date ON einv (company, inv_date);
            CREATE TABLE IF NOT EXISTS coverage (company TEXT, source TEXT, date_from TEXT, date_to TEXT,
                rows INTEGER, at TEXT, detail TEXT);
            CREATE TABLE IF NOT EXISTS taxpayers (tin TEXT PRIMARY KEY, name TEXT, address TEXT, town TEXT,
                at TEXT);
            CREATE TABLE IF NOT EXISTS einv_sales (company TEXT, eic TEXT, inv_no TEXT, inv_date TEXT,
                buyer_nipt TEXT, buyer_name TEXT, total REAL, currency TEXT, status TEXT, due TEXT, source TEXT,
                seen_at TEXT, PRIMARY KEY (company, eic));
            CREATE INDEX IF NOT EXISTS einv_sales_date ON einv_sales (company, inv_date);
        """)
    return _conn


def _read_state():
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _write_state(state):
    tmp = STATE_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(STATE_FILE)


def _log(entry):
    with _state_lock:
        state = _read_state()
        state["runs"] = ([entry] + state.get("runs", []))[:40]
        _write_state(state)


def settings():
    return store.get_settings().get("fiscal", {})


def nipt_keys(n):
    """Keys a supplier's NIPT is known by: upper case without spaces or signs; a foreign VAT number
    with a country prefix (DE123...) also without it."""
    k = re.sub(r"[^A-Z0-9]", "", str(n or "").upper())
    if not k:
        return []
    out = [k]
    if re.fullmatch(r"[A-Z]{2}\d{6,}", k):
        out.append(k[2:])
    return out


def _num(s, key, default):
    try:
        return float(s.get(key) if s.get(key) not in (None, "") else default)
    except (TypeError, ValueError):
        return float(default)


def _store_rows(company, rows, source, span=None):
    now = datetime.now().isoformat(timespec="seconds")
    with _lock:
        db = _db()
        db.executemany(
            "INSERT OR REPLACE INTO einv VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(company, r["eic"].lower(), r.get("iic"), r.get("inv_no"), r.get("date"),
              (r.get("seller_nipt") or "").strip().upper(), r.get("seller_name"),
              (r.get("buyer_nipt") or "").strip().upper(), r.get("total"), r.get("currency") or "ALL",
              r.get("doc_type"), r.get("status"), source, now) for r in rows if r.get("eic")])
        dates = sorted(r["date"] for r in rows if r.get("date"))
        a, b = span or ((dates[0], dates[-1]) if dates else (None, None))
        if a and b:
            db.execute("INSERT INTO coverage VALUES (?,?,?,?,?,?,?)", (company, source, a, b, len(rows), now, ""))
        db.commit()
    with _jobs_lock:
        for k in [k for k in _results if k[0] == company]:
            _results.pop(k, None)


# ---------------------------------------------------------------- Self-Care purchases book
def import_book(company, data, filename=""):
    """Reads the Self-Care purchases book (.xlsx) and keeps its invoices."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:                                     # noqa: BLE001
        raise ValueError("The file is not an Excel workbook (.xlsx).") from exc
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = None
    nipt = ""
    for i, r in enumerate(rows[:40]):
        cells = [str(c or "").strip().lower() for c in r]
        if cells and cells[0].startswith("nipt"):
            nipt = str(next((c for c in r[1:] if c), "") or "").strip().upper()
        if any(c.startswith("numri serial") for c in cells):
            head = i
            break
    if head is None:
        raise ValueError("This is not the Self-Care purchases book: no column 'Numri Serial'.")
    cells = [str(c or "").strip().lower() for c in rows[head]]

    def col(*names):
        for n in names:
            for j, c in enumerate(cells):
                if c.startswith(n):
                    return j
        return None
    ci = {"iic": col("nr fatur"), "eic": col("numri serial"), "date": col("data"),
          "name": col("emri tregtar"), "nipt": col("nip"), "total": None}
    ci["total"] = next((j for j, c in enumerate(str(x or "").lower() for x in rows[head - 1]) if c.startswith("totali")),
                       None) if head else None
    if ci["total"] is None:
        ci["total"] = 6
    out = []
    for r in rows[head + 1:]:
        eic = str(r[ci["eic"]] or "").strip() if ci["eic"] is not None and len(r) > ci["eic"] else ""
        if len(eic) != 36 or eic.count("-") != 4:
            continue
        d = r[ci["date"]]
        if isinstance(d, datetime):
            d = d.date().isoformat()
        elif isinstance(d, date):
            d = d.isoformat()
        else:
            s = str(d or "").strip()
            try:
                d = datetime.strptime(s[:10], "%d/%m/%Y").date().isoformat()
            except ValueError:
                d = s[:10] or None
        try:
            total = float(r[ci["total"]] or 0)
        except (TypeError, ValueError):
            total = None
        out.append({"eic": eic.lower(), "iic": str(r[ci["iic"]] or "").strip() if ci["iic"] is not None else None,
                    "date": d, "seller_name": str(r[ci["name"]] or "").strip() if ci["name"] is not None else "",
                    "seller_nipt": str(r[ci["nipt"]] or "").strip().upper() if ci["nipt"] is not None else "",
                    "buyer_nipt": nipt, "total": total, "currency": "ALL",
                    "doc_type": "credit note" if (total or 0) < 0 else "invoice"})
    if not out:
        raise ValueError("No invoice rows were found in the book.")
    own = (config.get("FISK_NIPT", "") or "").strip().upper()
    if own and nipt and nipt != own:
        raise ValueError(f"The book belongs to NIPT {nipt}, not to {own}.")
    dates = sorted(r["date"] for r in out if r["date"])
    year = month = None
    for r in rows[:12]:
        lab = str(r[0] or "").strip().lower()
        val = next((c for c in r[1:] if c not in (None, "")), None)
        if lab.startswith("viti"):
            year = val
        elif lab.startswith("muaji"):
            month = val
    span = None
    try:
        first = date(int(year), int(month), 1)
        last = (first.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        span = (first.isoformat(), min(last.isoformat(), dates[-1]) if dates else last.isoformat())
    except (TypeError, ValueError):
        pass
    _store_rows(company, out, "book", span)
    return {"rows": len(out), "from": dates[0] if dates else None, "to": dates[-1] if dates else None,
            "nipt": nipt, "file": filename}


# ---------------------------------------------------------------- the Tax Administration
def fetch(company, days=None, actor="scheduler", full=False):
    """Reads the e-invoices received by the platform from the e-invoice service. The first read
    (or a full one) covers the `days` setting; later ones start two days before the last read,
    since an invoice is listed by the date-time the platform received it."""
    days = int(days or _num(settings(), "days", 45))
    b = date.today()
    a = b - timedelta(days=days)
    with _state_lock:
        last = (_read_state().get("fetched_to") or {}).get(company)
    with _lock:
        held = _db().execute("SELECT COUNT(*) FROM einv WHERE company=? AND source='cis'", (company,)).fetchone()[0]
    if not held:
        full = True            # nothing from the service yet: read the whole window again
    if last and not full:
        a = max(a, date.fromisoformat(last) - timedelta(days=2))
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), "kind": "fetch", "company": company,
             "actor": actor, "from": a.isoformat(), "to": b.isoformat()}
    try:
        rows = cis_client.get_einvoices(a, b)
        read = len(rows)
        own = (config.get("FISK_NIPT", "") or "").strip().upper()
        for r in rows:
            if not r.get("buyer_nipt"):
                r["buyer_nipt"] = own
            r["doc_type"] = r.get("doc_type") or ("credit note" if (r.get("total") or 0) < 0 else "invoice")
        party = (config.get("FISK_PARTY_TYPE", "BUYER") or "BUYER").upper()
        if party == "BUYER" and own:
            rows = [r for r in rows if not r.get("buyer_nipt") or r["buyer_nipt"].upper() == own]
        _names(rows)
        _store_rows(company, rows, "cis", (a.isoformat(), b.isoformat()))
        with _state_lock:
            st = _read_state()
            st.setdefault("fetched_to", {})[company] = b.isoformat()
            _write_state(st)
        entry.update(status="ok", rows=len(rows), read=read, diag=dict(cis_client.last_diag))
        company_nipt = (settings().get("company_nipt") or "").strip().upper()
        cert_nipt = entry["diag"].get("cert_nipt") or own
        if company_nipt and cert_nipt and company_nipt != cert_nipt:
            entry["detail"] = (f"The certificate is for NIPT {cert_nipt}, not {company} ({company_nipt}): the service lists "
                               f"the e-invoices of {cert_nipt}.")
    except cis_client.NotConfigured as exc:
        entry.update(status="not configured", detail=str(exc))
    except Exception as exc:                                     # noqa: BLE001
        traceback.print_exc()
        entry.update(status="failed", detail=str(exc)[:400])
    _log(entry)
    store.audit("fiscal fetch", company, actor, {}, {k: entry.get(k) for k in ("status", "rows", "from", "to")},
                entry.get("detail", ""))
    return entry


def _names(rows, limit=40, nipt_key="seller_nipt", name_key="seller_name"):
    """Counterparty names: the service lists only the NIPT. Names already known (an imported book,
    an earlier lookup) are reused; up to `limit` new NIPTs a run are asked from the Tax
    Administration (GetTaxpayers)."""
    with _lock:
        db = _db()
        known = dict(db.execute("SELECT tin, name FROM taxpayers").fetchall())
        for tin, name in db.execute("SELECT seller_nipt, seller_name FROM einv WHERE seller_name<>'' "
                                    "GROUP BY seller_nipt").fetchall():
            known.setdefault(tin, name)
    asked = 0
    for r in rows:
        tin = (r.get(nipt_key) or "").upper()
        if r.get(name_key) or not tin:
            continue
        if tin not in known and asked < limit and re.fullmatch(r"[A-Z]\d{8}[A-Z]", tin):
            asked += 1
            try:
                tp = cis_client.get_taxpayer(tin) or {}
            except Exception as exc:                             # noqa: BLE001
                print(f"fiscal: taxpayer {tin} not read: {str(exc)[:120]}")
                tp = {}
            known[tin] = tp.get("name") or ""
            with _lock:
                _db().execute("INSERT OR REPLACE INTO taxpayers VALUES (?,?,?,?,?)",
                              (tin, tp.get("name") or "", tp.get("address") or "", tp.get("town") or "",
                               datetime.now().isoformat(timespec="seconds")))
                _db().commit()
        r[name_key] = known.get(tin) or ""


def coverage(company):
    """The loads and the periods they cover; `from`/`to` is the latest unbroken covered span."""
    with _lock:
        db = _db()
        cov = db.execute("SELECT source, date_from, date_to, rows, at FROM coverage WHERE company=? AND "
                         "source NOT LIKE '%:sales' ORDER BY at DESC", (company,)).fetchall()
        n = db.execute("SELECT COUNT(*) FROM einv WHERE company=?", (company,)).fetchone()[0]
    spans = sorted((c[1], c[2]) for c in cov if c[1] and c[2])
    merged = []
    for a, b in spans:
        if merged and a <= (date.fromisoformat(merged[-1][1]) + timedelta(days=1)).isoformat():
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    last = merged[-1] if merged else [None, None]
    return {"loads": [dict(zip(("source", "from", "to", "rows", "at"), c)) for c in cov[:12]],
            "spans": merged, "from": last[0], "to": last[1], "count": n}


# ---------------------------------------------------------------- Business Central
def _bc(client, company, a, b):
    """Posted purchase invoices, supplier ledger (lek amounts, credit memos), supplier NIPTs
    and purchase orders carrying an EIC."""
    start = (date.fromisoformat(a) - timedelta(days=60)).isoformat()
    inv = client._get_all(client._url(company, config.E_PINV), {
        "$filter": f"Posting_Date ge {start}",
        "$select": "No,Vendor_Invoice_No,Buy_from_Vendor_No,Buy_from_Vendor_Name,Pay_to_Vendor_No,Currency_Code,"
                   "Amount_Including_VAT,NIVF,IIC,Posting_Date,Document_Date,Cancelled,Corrective,Location_Code"})
    vle = client._get_all(client._url(company, config.E_VENDOR), {
        "$filter": f"Posting_Date ge {start} and (Document_Type eq 'Invoice' or Document_Type eq 'Credit Memo')",
        "$select": "Document_No,Document_Type,Vendor_No,Vendor_Name,Posting_Date,Document_Date,Amount_LCY,"
                   "Currency_Code"})
    vend = client._get_all(client._url(company, config.E_VENDOR_CARD),
                           {"$select": "No,Name,VAT_Registration_No"})
    try:
        po = client._get_all(client._url(company, config.E_PO), {
            "$filter": "NIVF ne ''",
            "$select": "No,Buy_from_Vendor_No,Buy_from_Vendor_Name,NIVF,Vendor_Invoice_No,Document_Date,Status"})
    except Exception as exc:                                     # noqa: BLE001
        print(f"fiscal: purchase orders not read: {str(exc)[:160]}")
        po = None
    return inv, vle, vend, po


def reconcile(einv, inv, vle, vend, po, a, b, today, s):
    tol = _num(s, "tolerance", 1)
    grace = int(_num(s, "grace_days", 7))
    nipt_of = {v["No"]: (v.get("VAT_Registration_No") or "").strip().upper() for v in vend}
    vend_of_nipt = {}
    for no, n in nipt_of.items():
        for key in nipt_keys(n):
            vend_of_nipt.setdefault(key, set()).add(no)

    def vendors(n):
        return set().union(*[vend_of_nipt.get(key, set()) for key in nipt_keys(n)] or [set()])
    lcy = {}
    cms = []
    for e in vle:
        if e.get("Document_Type") == "Invoice":
            lcy[e["Document_No"]] = abs(float(e.get("Amount_LCY") or 0))
        else:
            cms.append(e)

    def bc_amount(x):
        if x.get("Currency_Code"):
            return lcy.get(x["No"])
        return float(x.get("Amount_Including_VAT") or 0)

    def amount_for(x, cur):
        """BC amount in the e-invoice's currency: lek (posted lek amount) or the invoice's own currency."""
        if cur in ("", "ALL", "LEK"):
            return bc_amount(x)
        return float(x.get("Amount_Including_VAT") or 0) if (x.get("Currency_Code") or "").upper() == cur else None
    name_of_nipt = {}
    for v in vend:
        for key in nipt_keys(v.get("VAT_Registration_No")):
            name_of_nipt.setdefault(key, v.get("Name") or "")
    by_eic = {}
    for x in inv:
        k = (x.get("NIVF") or "").strip().lower()
        if k:
            by_eic.setdefault(k, []).append(x)
    by_vendor = {}
    for x in inv:
        by_vendor.setdefault(x.get("Buy_from_Vendor_No"), []).append(x)
    po_eic = {}
    for p in po or []:
        k = (p.get("NIVF") or "").strip().lower()
        if k:
            po_eic.setdefault(k, p)
    used, used_cm = set(), set()
    rows = []

    def near(d1, d2):
        try:
            return abs((date.fromisoformat(d1[:10]) - date.fromisoformat(d2[:10])).days)
        except (TypeError, ValueError):
            return 999
    for e in einv:
        total = e["total"] or 0.0
        cur = (e.get("currency") or "ALL").upper()
        cur = "ALL" if cur in ("", "LEK") else cur
        seller = e["seller_name"] or next((name_of_nipt[k] for k in nipt_keys(e["seller_nipt"]) if k in name_of_nipt), "")
        r = {"eic": e["eic"], "date": e["inv_date"], "seller": seller, "nipt": e["seller_nipt"] or "",
             "total": round(total, 2), "currency": cur, "source": e["source"], "bc_no": "", "bc_vendor": "",
             "bc_amount": None, "bc_date": "", "diff": None, "status": "", "note": "",
             "value": round(total, 2) if cur == "ALL" else 0.0}
        if (e.get("status") or "").upper() == "REFUSED":
            r.update(status="refused", note="refused by us at the Tax Administration")
            rows.append(r)
            continue
        if abs(total) < 0.005:
            r["status"] = "zero"
            rows.append(r)
            continue
        if total < 0:
            vs = vendors(r["nipt"])
            cand = [c for c in cms if c["Vendor_No"] in vs and (c["Document_No"], c["Vendor_No"]) not in used_cm
                    and cur == "ALL" and abs(abs(float(c.get("Amount_LCY") or 0)) - abs(total)) <= tol]
            cand.sort(key=lambda c: near(c.get("Document_Date") or c["Posting_Date"], r["date"] or ""))
            if cand:
                c = cand[0]
                used_cm.add((c["Document_No"], c["Vendor_No"]))
                r.update(status="credit memo", bc_no=c["Document_No"], bc_vendor=c.get("Vendor_Name") or "",
                         bc_amount=-abs(float(c.get("Amount_LCY") or 0)), bc_date=c["Posting_Date"])
            elif r["eic"] in po_eic:
                p = po_eic[r["eic"]]
                r.update(status="in order", bc_no=p["No"], bc_vendor=p.get("Buy_from_Vendor_Name") or "")
            else:
                r["status"] = "recent" if r["date"] and near(r["date"], today) <= grace else "missing credit note"
            rows.append(r)
            continue
        m = [x for x in by_eic.get(r["eic"], []) if x["No"] not in used]
        if m:
            x = m[0]
            used.add(x["No"])
            amt = amount_for(x, cur)
            r.update(bc_no=x["No"], bc_vendor=x.get("Buy_from_Vendor_Name") or "", bc_date=x.get("Posting_Date") or "",
                     bc_amount=round(amt, 2) if amt is not None else None)
            if cur != "ALL":
                r["value"] = round(bc_amount(x) or 0, 2)
            limit = max(tol, abs(total) * 0.01) if (x.get("Currency_Code") and cur == "ALL") else tol
            if amt is None:
                r.update(status="matched", note=f"fiscal in {cur}, BC in {x.get('Currency_Code') or 'ALL'}: "
                                                "amount not compared")
            elif abs(total - amt) <= limit:
                r["status"] = "matched"
            else:
                r.update(status="amount differs", diff=round(total - amt, 2),
                         note=f"BC in {x['Currency_Code']}" if x.get("Currency_Code") else "")
            rows.append(r)
            continue
        vs = vendors(r["nipt"])
        cand = [x for v in vs for x in by_vendor.get(v, []) if x["No"] not in used
                and not (x.get("NIVF") or "").strip() and (amount_for(x, cur) is not None)
                and abs(amount_for(x, cur) - total) <= tol and near(x.get("Document_Date") or "", r["date"] or "") <= 45]
        cand.sort(key=lambda x: near(x.get("Document_Date") or "", r["date"] or ""))
        if cand:
            x = cand[0]
            used.add(x["No"])
            r.update(status="matched, no EIC in BC", bc_no=x["No"], bc_vendor=x.get("Buy_from_Vendor_Name") or "",
                     bc_date=x.get("Posting_Date") or "", bc_amount=round(amount_for(x, cur), 2))
            if cur != "ALL":
                r["value"] = round(bc_amount(x) or 0, 2)
        elif cur == "ALL" and r["source"] == "cis" and [c for c in cms if c["Vendor_No"] in vs and (c["Document_No"], c["Vendor_No"]) not in used_cm
                                                  and abs(abs(float(c.get("Amount_LCY") or 0)) - total) <= tol]:
            # the service lists credit notes with a positive amount: a posted credit memo of the supplier
            c = sorted([c for c in cms if c["Vendor_No"] in vs and (c["Document_No"], c["Vendor_No"]) not in used_cm
                        and abs(abs(float(c.get("Amount_LCY") or 0)) - total) <= tol],
                       key=lambda c: near(c.get("Document_Date") or c["Posting_Date"], r["date"] or ""))[0]
            used_cm.add((c["Document_No"], c["Vendor_No"]))
            r.update(status="credit memo", bc_no=c["Document_No"], bc_vendor=c.get("Vendor_Name") or "",
                     bc_amount=-abs(float(c.get("Amount_LCY") or 0)), bc_date=c["Posting_Date"], value=-total)
        elif r["eic"] in po_eic:
            p = po_eic[r["eic"]]
            r.update(status="in order", bc_no=p["No"], bc_vendor=p.get("Buy_from_Vendor_Name") or "",
                     note=f"purchase order {p['No']} ({p.get('Status') or ''}), invoice not posted")
        elif r["date"] and near(r["date"], today) <= grace:
            r["status"] = "recent"
        else:
            r["status"] = "missing"
            if not vs:
                r["note"] = "no supplier in BC carries this NIPT"
        rows.append(r)

    # ---- from the BC side
    eics = {e["eic"] for e in einv}
    in_period = [x for x in inv if a <= (x.get("Document_Date") or "")[:10] <= b]
    dup = []
    for k, xs in by_eic.items():
        if len(xs) > 1:
            same = len({round(bc_amount(x) or 0, 2) for x in xs}) == 1
            for x in xs:
                dup.append({"eic": k, "bc_no": x["No"], "vendor": x.get("Buy_from_Vendor_Name") or "",
                            "posting_date": x.get("Posting_Date"), "document_date": x.get("Document_Date"),
                            "amount": round(bc_amount(x) or 0, 2), "currency": x.get("Currency_Code") or "",
                            "kind": "same amount" if same else "different amounts", "in_fiscal": k in eics})
    dup.sort(key=lambda d: (d["kind"] != "same amount", -d["amount"], d["eic"]))
    not_fiscal = [{"bc_no": x["No"], "vendor": x.get("Buy_from_Vendor_Name") or "", "eic": x["NIVF"],
                   "document_date": x.get("Document_Date"), "amount": round(bc_amount(x) or 0, 2),
                   "nipt": nipt_of.get(x.get("Buy_from_Vendor_No"), "")}
                  for x in in_period if (x.get("NIVF") or "").strip() and x["NIVF"].strip().lower() not in eics]
    no_eic = [{"bc_no": x["No"], "vendor": x.get("Buy_from_Vendor_Name") or "", "vendor_no": x.get("Buy_from_Vendor_No"),
               "vendor_invoice": x.get("Vendor_Invoice_No") or "", "document_date": x.get("Document_Date"),
               "amount": round(bc_amount(x) or 0, 2), "nipt": nipt_of.get(x.get("Buy_from_Vendor_No"), ""),
               "location": x.get("Location_Code") or ""}
              for x in in_period if not (x.get("NIVF") or "").strip() and x["No"] not in used]
    no_eic.sort(key=lambda r: -r["amount"])
    not_fiscal.sort(key=lambda r: -r["amount"])

    def tally(st):
        sel = [r for r in rows if r["status"] == st]
        return [len(sel), round(sum(r["value"] for r in sel), 0)]
    statuses = ["matched", "matched, no EIC in BC", "amount differs", "credit memo", "in order", "recent",
                "missing", "missing credit note", "zero", "refused"]
    k = {st: tally(st) for st in statuses}
    dup_same = [d for d in dup if d["kind"] == "same amount"]
    groups = {}
    for d in dup_same:
        groups.setdefault(d["eic"], []).append(d)
    k["posted twice"] = [len(groups), round(sum(g[0]["amount"] * (len(g) - 1) for g in groups.values()), 0)]
    k["bc not in fiscal"] = [len(not_fiscal), round(sum(r["amount"] for r in not_fiscal), 0)]
    k["bc without eic"] = [len(no_eic), round(sum(r["amount"] for r in no_eic), 0)]
    k["fiscal"] = [len(rows), round(sum(r["value"] for r in rows), 0)]
    k["foreign"] = [sum(1 for r in rows if r["currency"] != "ALL"),
                    sorted({r["currency"] for r in rows if r["currency"] != "ALL"})]
    k["amount differs net"] = round(sum(r["diff"] or 0 for r in rows if r["status"] == "amount differs"), 0)
    miss = [r for r in rows if r["status"] in ("missing", "missing credit note")]
    by_sup = {}
    for r in miss:
        t = by_sup.setdefault(r["nipt"] or r["seller"], {"seller": r["seller"], "nipt": r["nipt"], "count": 0,
                                                          "value": 0.0, "oldest": r["date"]})
        t["count"] += 1
        t["value"] += r["value"]
        t["oldest"] = min(t["oldest"] or r["date"], r["date"] or t["oldest"])
    order = {st: i for i, st in enumerate(["missing", "amount differs", "missing credit note", "in order",
                                           "matched, no EIC in BC", "recent", "credit memo", "refused", "zero",
                                           "matched"])}
    rows.sort(key=lambda r: (order.get(r["status"], 99), -abs(r["total"])))
    return {"kpis": k, "rows": rows, "posted_twice": dup, "bc_not_in_fiscal": not_fiscal, "bc_without_eic": no_eic,
            "missing_by_supplier": sorted(({**v, "value": round(v["value"], 0)} for v in by_sup.values()),
                                          key=lambda v: -v["value"])[:100],
            "orders_read": po is not None}


# ---------------------------------------------------------------- the report (background)
def request(client, company, date_from=None, date_to=None, refresh=False):
    s = settings()
    cov = coverage(company)
    if not cov["count"]:
        return {"available": False, "coverage": cov, "certificate": cis_client.certificate_info(),
                "error": "No fiscalized invoices are held yet: read them from the Tax Administration or import "
                         "the Self-Care purchases book."}
    b = (date_to or cov["to"] or date.today().isoformat())[:10]
    a = (date_from or max(cov["from"], (date.fromisoformat(b) - timedelta(days=int(_num(s, "days", 45)))).isoformat()))[:10]
    key = (company, a, b, cov["count"], cov["loads"][0]["at"] if cov["loads"] else "")
    with _jobs_lock:
        cached = _results.get(key)
        job = _jobs.get(company)
        running = job and job["state"] == "running"
    if cached and not refresh and time.time() - cached[0] < 1800:
        return cached[1]
    if not running:
        with _jobs_lock:
            job = _jobs[company] = {"state": "running", "stage": "Business Central", "started": time.time()}
        threading.Thread(target=_run, args=(client, company, a, b, key, job), daemon=True, name="fiscal").start()
    return {"pending": True, "progress": {"stage": job.get("stage")}, "from": a, "to": b,
            "note": "Reading the posted purchase invoices, the supplier ledger and the suppliers from Business Central."}


def _run(client, company, a, b, key, job):
    try:
        inv, vle, vend, po = _bc(client, company, a, b)
        with _jobs_lock:
            job["stage"] = "Matching"
        with _lock:
            cur = _db().execute("SELECT eic, iic, inv_no, inv_date, seller_nipt, seller_name, total, currency, "
                                "doc_type, status, source FROM einv WHERE company=? AND inv_date>=? AND inv_date<=?",
                                (company, a, b))
            cols = [c[0] for c in cur.description]
            einv = [dict(zip(cols, r)) for r in cur.fetchall()]
        out = reconcile(einv, inv, vle, vend, po, a, b, date.today().isoformat(), settings())
        out.update(company=company, date_from=a, date_to=b, checked_at=datetime.now().isoformat(timespec="seconds"),
                   coverage=coverage(company), certificate=cis_client.certificate_info(), definitions=DEFINITIONS,
                   bc_read={"invoices": len(inv), "ledger": len(vle), "suppliers": len(vend),
                            "orders": len(po) if po is not None else None})
        with _jobs_lock:
            _results[key] = (time.time(), out)
            for k, _ in sorted(_results.items(), key=lambda kv: kv[1][0])[:-4]:
                _results.pop(k, None)
            job["state"] = "done"
    except Exception as exc:                                     # noqa: BLE001
        traceback.print_exc()
        with _jobs_lock:
            _results[key] = (time.time(), {"available": False, "error": f"Business Central could not be read: "
                                                                         f"{str(exc)[:300]}"})
            job["state"] = "failed"


def result_now(client, company, timeout=600):
    """Waits for the report (the daily email runs outside any request)."""
    t0 = time.time()
    r = request(client, company)
    while r.get("pending") and time.time() - t0 < timeout:
        time.sleep(3)
        r = request(client, company)
    return r


DEFINITIONS = {
    "source": "Fiscalized invoices: the e-invoice service of the Tax Administration (read with the fiscalization "
              "certificate) or the Self-Care purchases book imported by hand. Business Central is only read.",
    "matched": "Matched: the e-invoice code (EIC) is the NIVF of a posted purchase invoice and the amounts agree "
               "within the tolerance (foreign-currency invoices at the lek amount BC posted).",
    "matched_no_eic": "Matched, no EIC in BC: same supplier NIPT and amount on a posted invoice without NIVF; the "
                      "NIVF should be filled in.",
    "amount": "Amount differs: the EIC is posted with a different amount.",
    "in_order": "In order: the EIC is on a purchase order whose invoice is not posted yet.",
    "missing": "Missing: fiscalized to us, older than the grace days, and not found in BC by EIC, by NIPT and amount, "
               "or on a purchase order. Recent: the same, but still within the grace days.",
    "credit": "Credit notes (negative totals) are matched to posted credit memos by supplier NIPT and amount.",
    "dates": "Date: the invoice date in the Self-Care book; the date-time the platform received the invoice when it "
             "comes from the e-invoice service. Amounts in another currency are compared in that currency with the "
             "BC invoice in the same currency; values in the summary are in lek (BC's posted lek amount).",
    "refused": "Refused: invoices we refused at the Tax Administration are listed apart and not counted as missing.",
    "posted_twice": "Posted twice: the same EIC on two or more posted purchase invoices with the same amount. Check "
                    "whether one was reversed before acting.",
    "bc_side": "BC not in fiscal: posted invoices of the period whose EIC the Tax Administration does not list for us. "
               "BC without EIC: posted invoices of the period without NIVF and not matched by NIPT and amount "
               "(often farmers and individuals who do not issue e-invoices).",
}


# ---------------------------------------------------------------- Excel
def render_xlsx(r):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    fill, bold = PatternFill("solid", fgColor="E4F3F9"), Font(bold=True)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = fill, bold
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    k = r["kpis"]
    ws = wb.active
    ws.title = "Summary"
    labels = [("fiscal", "Fiscalized invoices in the period"), ("matched", "Matched"),
              ("matched, no EIC in BC", "Matched by NIPT and amount, no EIC in BC"),
              ("amount differs", "EIC posted with a different amount"), ("credit memo", "Credit notes matched"),
              ("in order", "On a purchase order, not posted"), ("recent", "Recent, within the grace days"),
              ("missing", "Missing in BC"), ("missing credit note", "Credit notes missing in BC"),
              ("zero", "Zero value"), ("refused", "Refused by us at the Tax Administration"), ("posted twice", "BC: same EIC posted twice, same amount (extra value)"),
              ("bc not in fiscal", "BC: EIC not listed by the Tax Administration"),
              ("bc without eic", "BC: posted invoices without EIC")]
    sheet(ws, ["Check", "Count", "Value (ALL)"], [[l, k[key][0], k[key][1]] for key, l in labels], [52, 10, 16])
    ws.append([])
    ws.append([f"{r['company']} · invoices dated {r['date_from']} to {r['date_to']} · checked {r['checked_at']}"])
    cols = ["Status", "Date", "Supplier", "NIPT", "Total", "Currency", "EIC", "BC document", "BC supplier",
            "BC posting date", "BC amount", "Difference", "Note", "Source"]
    sheet(wb.create_sheet("Fiscalized invoices"), cols,
          [[x["status"], x["date"], x["seller"], x["nipt"], x["total"], x.get("currency", "ALL"), x["eic"], x["bc_no"],
            x["bc_vendor"], x["bc_date"], x["bc_amount"], x["diff"], x["note"],
            "Tax Administration" if x["source"] == "cis" else "Self-Care book"] for x in r["rows"]],
          [20, 11, 30, 13, 14, 8, 38, 11, 26, 12, 14, 12, 34, 16])
    sheet(wb.create_sheet("Missing by supplier"), ["Supplier", "NIPT", "Invoices", "Value (ALL)", "Oldest"],
          [[x["seller"], x["nipt"], x["count"], x["value"], x["oldest"]] for x in r["missing_by_supplier"]],
          [34, 13, 10, 14, 11])
    sheet(wb.create_sheet("Posted twice"), ["EIC", "BC document", "Supplier", "Posting date", "Document date",
                                            "Amount (ALL)", "Currency", "Kind", "Listed by Tax Administration"],
          [[d["eic"], d["bc_no"], d["vendor"], d["posting_date"], d["document_date"], d["amount"], d["currency"],
            d["kind"], "yes" if d["in_fiscal"] else "no"] for d in r["posted_twice"]],
          [38, 11, 28, 12, 12, 14, 9, 16, 12])
    sheet(wb.create_sheet("BC not in fiscal"), ["BC document", "Supplier", "NIPT", "EIC", "Document date", "Amount"],
          [[x["bc_no"], x["vendor"], x["nipt"], x["eic"], x["document_date"], x["amount"]] for x in r["bc_not_in_fiscal"]],
          [11, 30, 13, 38, 12, 14])
    sheet(wb.create_sheet("BC without EIC"), ["BC document", "Supplier", "Supplier no.", "NIPT", "Supplier invoice",
                                              "Document date", "Amount", "Location"],
          [[x["bc_no"], x["vendor"], x["vendor_no"], x["nipt"], x["vendor_invoice"], x["document_date"], x["amount"],
            x["location"]] for x in r["bc_without_eic"]], [11, 30, 11, 13, 16, 12, 14, 12])
    ws = wb.create_sheet("Definitions")
    for key, v in DEFINITIONS.items():
        ws.append([key, v])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_name(r):
    return f"RHG_{r['company'].replace(' ', '_')}_fiscal_vs_BC_{r['date_from']}_{r['date_to']}.xlsx"


# ---------------------------------------------------------------- the email
def _m(v):
    return f"{v:,.0f}".replace(",", " ")


def render_email(r, test=False):
    import html as H
    k = r["kpis"]
    url = f"{config.PUBLIC_URL}/finance/#fiscal"
    rows = [("Missing in BC", k["missing"], True), ("Credit notes missing in BC", k["missing credit note"], True),
            ("EIC posted with a different amount", k["amount differs"], True),
            ("Same EIC posted twice (extra value)", k["posted twice"], True),
            ("On a purchase order, not posted", k["in order"], False),
            ("Matched by NIPT and amount, NIVF to fill in", k["matched, no EIC in BC"], False),
            ("Recent, within the grace days", k["recent"], False), ("Matched", k["matched"], False)]
    trs = "".join(f'<tr><td style="padding:6px 10px;border-bottom:1px solid #e6e9ea">{H.escape(l)}</td>'
                  f'<td style="padding:6px 10px;border-bottom:1px solid #e6e9ea;text-align:right;'
                  f'{"color:#c0392b;font-weight:700" if warn and v[0] else ""}">{v[0]}</td>'
                  f'<td style="padding:6px 10px;border-bottom:1px solid #e6e9ea;text-align:right">{_m(v[1])}</td></tr>'
                  for l, v, warn in rows)
    sup = "".join(f'<tr><td style="padding:5px 10px">{H.escape(x["seller"])}</td><td style="padding:5px 10px">'
                  f'{H.escape(x["nipt"])}</td><td style="padding:5px 10px;text-align:right">{x["count"]}</td>'
                  f'<td style="padding:5px 10px;text-align:right">{_m(x["value"])}</td>'
                  f'<td style="padding:5px 10px">{H.escape(x["oldest"] or "")}</td></tr>'
                  for x in r["missing_by_supplier"][:15])
    html_body = f"""<div style="font-family:Segoe UI,Arial,sans-serif;color:#0e2733;max-width:760px">
<h2 style="margin:0 0 4px">{'[TEST] ' if test else ''}Fiscalized purchase invoices against Business Central</h2>
<p style="color:#5c7482;margin:0 0 14px">{H.escape(r['company'])} · invoices dated {r['date_from']} to {r['date_to']} ·
{k['fiscal'][0]} invoices, {_m(k['fiscal'][1])} ALL</p>
<table style="border-collapse:collapse;width:100%;font-size:14px"><tr style="background:#e4f3f9">
<th style="padding:6px 10px;text-align:left">Check</th><th style="padding:6px 10px;text-align:right">Count</th>
<th style="padding:6px 10px;text-align:right">Value (ALL)</th></tr>{trs}</table>
{f'<h3 style="margin:18px 0 6px">Missing in BC, by supplier</h3><table style="border-collapse:collapse;width:100%;font-size:13px"><tr style="background:#e4f3f9"><th style="padding:5px 10px;text-align:left">Supplier</th><th style="padding:5px 10px;text-align:left">NIPT</th><th style="padding:5px 10px;text-align:right">Invoices</th><th style="padding:5px 10px;text-align:right">Value</th><th style="padding:5px 10px;text-align:left">Oldest</th></tr>{sup}</table>' if sup else ''}
<p style="margin:16px 0 0">Every line is in the attached Excel file and on the screen: <a href="{url}">{url}</a></p>
<p style="color:#5c7482;font-size:12px">Business Central is only read. Source of the fiscalized invoices:
{'the Tax Administration e-invoice service' if any(x["source"] == "cis" for x in r["rows"]) else 'the Self-Care purchases book'}.</p></div>"""
    text = (f"Fiscalized purchase invoices against BC, {r['company']}, {r['date_from']} to {r['date_to']}\n"
            + "\n".join(f"{l}: {v[0]} ({_m(v[1])} ALL)" for l, v, _ in rows) + f"\n\n{url}\n")
    subj = (f"{'[TEST] ' if test else ''}Fiscal vs BC · {r['company']} · {k['missing'][0]} missing, "
            f"{k['amount differs'][0]} amount differences, {k['posted twice'][0]} posted twice")
    return subj, html_body, text


def send(client, company, test=False, only_to=None, actor="scheduler"):
    with _run_lock:
        s = settings()
        entry = {"ts": datetime.now().isoformat(timespec="seconds"), "kind": "email", "company": company,
                 "test": bool(test), "actor": actor}
        try:
            to, _ = arap_notify.parse_emails(s.get("to"))
            cc, _ = arap_notify.parse_emails(s.get("cc"))
            if only_to:
                to, cc = only_to, []
            r = result_now(client, company)
            if r.get("available") is False or r.get("pending"):
                raise RuntimeError(r.get("error") or "The report is not ready.")
            k = r["kpis"]
            issues = k["missing"][0] + k["missing credit note"][0] + k["amount differs"][0] + k["posted twice"][0]
            # the sales side, when e-invoices issued by us are held
            import fiscal_sales
            sales, s_issues, s_html, s_text = None, 0, "", ""
            try:
                if fiscal_sales.coverage(company)["loads"]:
                    sr = fiscal_sales.result_now(client, company)
                    if sr.get("kpis"):
                        sales = sr
                        s_issues, s_html, s_text = fiscal_sales.email_block(sr)
            except Exception as exc:                             # noqa: BLE001
                s_html = f'<p style="color:#b26a00">The sales side could not be read: {str(exc)[:200]}</p>'
            issues += s_issues
            entry.update(issues=issues, to=to, cc=cc)
            if not to:
                entry.update(status="skipped", detail="No recipient is set in Settings.")
            elif not issues and not s.get("send_empty") and not test:
                entry.update(status="skipped", detail="Nothing to report.")
            else:
                subj, html_body, text = render_email(r, test)
                xl = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                files = [(xlsx_name(r), xl, render_xlsx(r))]
                if sales:
                    files.append((fiscal_sales.xlsx_name(sales), xl, fiscal_sales.render_xlsx(sales)))
                    subj += f" · sales: {sales['kpis']['missing'][0]} missing, {sales['kpis']['b2b without eic'][0]} " \
                            "B2B without e-invoice"
                if s_html:
                    html_body = html_body.replace('<p style="margin:16px 0 0">Every line', s_html +
                                                  '<p style="margin:16px 0 0">Every line', 1)
                    text += s_text
                arap_notify.send_mail(to, [c for c in cc if c not in to], subj, html_body, text, files)
                entry.update(status="sent", detail=subj)
        except Exception as exc:                                 # noqa: BLE001
            traceback.print_exc()
            entry.update(status="failed", detail=str(exc)[:400])
        _log(entry)
        store.audit("email", f"fiscal vs BC · {company}", actor, {},
                    {k2: entry.get(k2) for k2 in ("status", "issues", "to", "test")}, entry.get("detail", ""))
        return entry


# ---------------------------------------------------------------- jobs started from the screen
def start(kind, client, company, actor, **kw):
    with _jobs_lock:
        jid = uuid.uuid4().hex
        job = {"id": jid, "kind": kind, "actor": actor, "state": "running", "result": None,
               "started": datetime.now().isoformat(timespec="seconds")}
        _jobs["ui-" + jid] = job

    def run():
        try:
            if kind == "fetch":
                res = fetch(company, actor=actor)
            elif kind == "fetch sales":
                import fiscal_sales
                res = fiscal_sales.fetch(company, actor=actor)
            else:
                res = send(client, company, test=True, only_to=kw.get("to"), actor=actor)
        except Exception as exc:                                 # noqa: BLE001
            res = {"status": "failed", "detail": str(exc)[:300]}
        with _jobs_lock:
            job.update(state="done", result=res)
    threading.Thread(target=run, daemon=True, name=f"fiscal-{kind}").start()
    return dict(job)


def job(jid):
    with _jobs_lock:
        j = _jobs.get("ui-" + jid)
        return dict(j) if j else None


def status(company):
    with _state_lock:
        st = _read_state()
    import fiscal_sales
    return {"settings": settings(), "certificate": cis_client.certificate_info(), "coverage": coverage(company),
            "coverage_sales": fiscal_sales.coverage(company),
            "runs": st.get("runs", [])[:15], "next_run": next_run(), "mail": arap_notify.mail_status()}


# ---------------------------------------------------------------- the daily run
def _target(now, s):
    hh, _, mm = (s.get("time") or "07:30").partition(":")
    try:
        return now.replace(hour=int(hh), minute=int(mm or 0), second=0, microsecond=0)
    except ValueError:
        return None


def next_run(now=None):
    s = settings()
    if not s.get("enabled"):
        return None
    now = now or datetime.now()
    t = _target(now, s)
    if t is None:
        return None
    with _state_lock:
        last = _read_state().get("last_daily")
    if last == now.date().isoformat() or t + timedelta(hours=6) < now:
        t += timedelta(days=1)
    while s.get("working_days_only", True) and t.weekday() >= 5:
        t += timedelta(days=1)
    return t.isoformat(timespec="minutes")


def _due(now):
    s = settings()
    if not s.get("enabled"):
        return False
    if s.get("working_days_only", True) and now.weekday() >= 5:
        return False
    t = _target(now, s)
    if t is None or now < t or now > t + timedelta(hours=6):
        return False
    with _state_lock:
        st = _read_state()
    today = now.date().isoformat()
    if st.get("last_daily") == today:
        return False
    tries = st.get("tries", {})
    return not (tries.get("day") == today and (tries.get("count", 0) >= 3 or time.time() - tries.get("at", 0) < 900))


def run_daily(client, company):
    import fiscal_sales
    f = fetch(company)
    if f["status"] == "ok":
        fiscal_sales.fetch(company)
    e = send(client, company) if f["status"] in ("ok", "not configured") else None
    ok = f["status"] in ("ok", "not configured") and e and e["status"] != "failed"
    today = date.today().isoformat()
    with _state_lock:
        st = _read_state()
        if ok:
            st["last_daily"] = today
            st.pop("tries", None)
        else:
            tries = st.get("tries", {})
            count = tries.get("count", 0) + 1 if tries.get("day") == today else 1
            st["tries"] = {"day": today, "count": count, "at": time.time()}
            if count >= 3:
                st["last_daily"] = today
        _write_state(st)
    return f, e


def loop(client, company):
    while True:
        try:
            if _due(datetime.now()):
                run_daily(client, company)
        except Exception:                                        # noqa: BLE001
            traceback.print_exc()
        time.sleep(60)


def start_scheduler(client, company):
    threading.Thread(target=loop, args=(client, company), daemon=True, name="fiscal-daily").start()
